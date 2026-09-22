"""导出命令回归 —— 不连 SAP、不碰数据库。

三个 SAP 取数入口（_query_master / _query_stock / _query_price）全部 mock，
纯函数部分（行集合、透视、写出）直接单测。

命令本身不 import 任何 Django 模型，所以整个模块是 SimpleTestCase，
测试跑得很快。
"""

from __future__ import annotations

import io
import tempfile
from pathlib import Path
from unittest.mock import patch

import polars as pl
from openpyxl import load_workbook

from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import SimpleTestCase

from app_sap_services.exceptions import SAPBusinessError, SAPRfcError
from app_sap_services.schemas.messages import SapMessage
from app_sap_services.management.commands import export_material_price_stock as mod
from app_sap_services.management.commands.export_material_price_stock import (
    Command,
    build_data_rows,
    build_rows,
    fill_forward,
    write_workbook,
)


HEALTH = {"status": "healthy", "ashost": "192.168.103.182", "client": "800"}
HEALTH_PATCH = (
    "app_sap_services.management.commands.export_material_price_stock.sap_health_check"
)


# ----------------------------------------------------------------------
# 帧构造
# ----------------------------------------------------------------------

MASTER_SCHEMA = {
    "MATNR": pl.Utf8, "WERKS": pl.Utf8, "MAKTX": pl.Utf8, "MEINS": pl.Utf8,
    "NORMT": pl.Utf8, "GROES": pl.Utf8, "MTART": pl.Utf8, "MATKL": pl.Utf8,
    "LVORM": pl.Utf8, "LVORMC": pl.Utf8, "ZZTEXT1": pl.Utf8, "ZZFIGURE_NO": pl.Utf8,
}

STOCK_SCHEMA = {
    "MATNR": pl.Utf8, "MAKTX": pl.Utf8, "WERKS": pl.Utf8, "LGORT": pl.Utf8,
    "CHARG": pl.Utf8, "CLABS": pl.Float64, "EISBE": pl.Float64,
}

PRICE_SCHEMA = {
    "KALNR": pl.Utf8, "BDATJ": pl.Int64, "POPER": pl.Int64, "PEINH": pl.Int64,
    "VPRSV": pl.Utf8, "STPRS": pl.Float64, "PVPRS": pl.Float64, "WAERS": pl.Utf8,
    "SALK3": pl.Float64, "SALKV": pl.Float64, "MATNR": pl.Utf8, "BWKEY": pl.Utf8,
    "VERPR": pl.Float64,
}


def master_row(matnr, werks, maktx="原料X", meins="KG", lvorm=""):
    return {
        "MATNR": matnr, "WERKS": werks, "MAKTX": maktx, "MEINS": meins,
        "NORMT": "", "GROES": "", "MTART": "ROH", "MATKL": "",
        "LVORM": lvorm, "LVORMC": "", "ZZTEXT1": "", "ZZFIGURE_NO": "",
    }


def master_df(rows) -> pl.DataFrame:
    if not rows:
        return pl.DataFrame(schema=MASTER_SCHEMA)
    return pl.DataFrame(rows, schema=MASTER_SCHEMA)


def stock_row(matnr, werks, clabs=10.0, loc="0001", charg="B1", maktx=""):
    return {
        "MATNR": matnr, "MAKTX": maktx, "WERKS": werks, "LGORT": loc,
        "CHARG": charg, "CLABS": clabs, "EISBE": 1.0,
    }


def stock_df(rows) -> pl.DataFrame:
    if not rows:
        return pl.DataFrame(schema=STOCK_SCHEMA)
    return pl.DataFrame(rows, schema=STOCK_SCHEMA)


def price_row(matnr, bwkey="3011", bdatj=2026, poper=9, peinh=1, verpr=12.5,
              waers="CNY", vprsv="V", kalnr="000000000001"):
    return {
        "KALNR": kalnr, "BDATJ": bdatj, "POPER": poper, "PEINH": peinh,
        "VPRSV": vprsv, "STPRS": 0.0, "PVPRS": 0.0, "WAERS": waers,
        "SALK3": 0.0, "SALKV": 0.0, "MATNR": matnr, "BWKEY": bwkey,
        "VERPR": verpr,
    }


def price_df(rows) -> pl.DataFrame:
    if not rows:
        return pl.DataFrame(schema=PRICE_SCHEMA)
    return pl.DataFrame(rows, schema=PRICE_SCHEMA)


# ======================================================================
# 纯函数：行集合
# ======================================================================

class BuildRowsTests(SimpleTestCase):

    def test_union_of_all_three_sources(self):
        master = {("A", "3011")}
        stock = {("B", "3011"): 1.0}
        price = {("C", "3011"): {202601: 1.0}}
        self.assertEqual(
            build_rows(master, stock, price),
            [("A", "3011"), ("B", "3011"), ("C", "3011")],
        )

    def test_master_rows_without_any_data_are_kept(self):
        """默认忠于「物料清单来自主数据」—— 无数据行也保留。"""
        rows = build_rows({("A", "3011")}, {}, {})
        self.assertEqual(rows, [("A", "3011")])

    def test_only_with_data_drops_empty_rows(self):
        rows = build_rows(
            {("A", "3011"), ("B", "3011")},
            {("B", "3011"): 1.0}, {},
            only_with_data=True,
        )
        self.assertEqual(rows, [("B", "3011")])

    def test_sorted_by_material_then_plant(self):
        rows = build_rows({("B", "3011"), ("A", "3020"), ("A", "3011")}, {}, {})
        self.assertEqual(rows, [("A", "3011"), ("A", "3020"), ("B", "3011")])


class FillForwardTests(SimpleTestCase):

    def test_fills_gaps_between_known_months(self):
        months = [202601, 202602, 202603]
        out = fill_forward({202601: 5.0, 202603: 9.0}, months)
        self.assertEqual(out, {202601: 5.0, 202602: 5.0, 202603: 9.0})

    def test_does_not_backfill_before_first_known_month(self):
        """窗口开始前没有数据，不该凭空造出更早的价格。"""
        months = [202601, 202602, 202603]
        out = fill_forward({202603: 9.0}, months)
        self.assertEqual(out, {202603: 9.0})

    def test_empty_input(self):
        self.assertEqual(fill_forward({}, [202601, 202602]), {})


class BuildDataRowsTests(SimpleTestCase):

    def test_missing_month_is_blank_not_zero(self):
        """空白 = 该期间无估值变动记录，与「价格为 0」是两回事。"""
        rows = build_data_rows(
            [("A", "3011")], [202601, 202602],
            {"A": "原料"}, {"A": "KG"},
            {}, {("A", "3011"): {202601: 3.5}},
        )
        self.assertEqual(rows[0], ["A", "3011", "原料", 0.0, 3.5, None, "KG"])

    def test_missing_stock_is_zero(self):
        rows = build_data_rows(
            [("A", "3011")], [202601], {}, {}, {}, {})
        self.assertEqual(rows[0][3], 0.0)

    def test_stock_and_metadata_are_looked_up(self):
        rows = build_data_rows(
            [("A", "3011")], [202601],
            {"A": "原料"}, {"A": "L"},
            {("A", "3011"): 7.25}, {},
        )
        self.assertEqual(rows[0], ["A", "3011", "原料", 7.25, None, "L"])


# ======================================================================
# 纯函数：写出
# ======================================================================

class WriteWorkbookTests(SimpleTestCase):

    def _write(self, data_rows, month_labels=("2025-10", "2025-11"), notes=()):
        buf = io.BytesIO()
        write_workbook(buf, list(month_labels), data_rows, notes)
        buf.seek(0)
        return load_workbook(buf)

    def test_headers_and_month_columns(self):
        wb = self._write([["A", "3011", "原料", 1.0, 2.5, None, "KG"]])
        ws = wb[mod.DATA_SHEET]
        self.assertEqual(
            [c.value for c in ws[1]],
            ["原材料SAP编码", "工厂", "原材料名称", "原材料实时库存",
             "2025-10", "2025-11", "基本计量单位"],
        )
        self.assertEqual(wb.sheetnames[0], mod.DATA_SHEET)

    def test_blank_price_cell_vs_zero_stock_cell(self):
        wb = self._write([["A", "3011", "原料", 0.0, None, None, "KG"]])
        ws = wb[mod.DATA_SHEET]
        self.assertEqual(ws["D2"].value, 0.0)
        self.assertIsNone(ws["E2"].value)
        self.assertIsNone(ws["F2"].value)

    def test_number_formats(self):
        wb = self._write([["A", "3011", "原料", 1.5, 2.5, 3.5, "KG"]])
        ws = wb[mod.DATA_SHEET]
        self.assertEqual(ws["A2"].number_format, "@")
        self.assertEqual(ws["B2"].number_format, "@")
        self.assertEqual(ws["D2"].number_format, "0.000")
        self.assertEqual(ws["E2"].number_format, "0.0000")
        self.assertEqual(ws["F2"].number_format, "0.0000")

    def test_freeze_panes_keeps_code_and_stock_visible(self):
        """月份列多，冻 A2 横滚就看不见编码了。"""
        wb = self._write([["A", "3011", "原料", 1.0, 2.5, None, "KG"]])
        self.assertEqual(wb[mod.DATA_SHEET].freeze_panes, "E2")

    def test_sap_code_is_stored_as_text(self):
        wb = self._write([["1001000001", "3011", "原料", 1.0, None, None, "KG"]])
        cell = wb[mod.DATA_SHEET]["A2"]
        self.assertEqual(cell.value, "1001000001")
        self.assertEqual(cell.data_type, "s")

    def test_formula_like_name_is_forced_to_text(self):
        """SAP 物料描述是人填的，'=' 开头会被当成公式。"""
        wb = self._write([["A", "3011", "=1+1", 1.0, None, None, "KG"]])
        cell = wb[mod.DATA_SHEET]["C2"]
        self.assertEqual(cell.value, "=1+1")
        self.assertEqual(cell.data_type, "s")

    def test_notes_sheet(self):
        wb = self._write([], notes=[("导出时间", "2026-09-21"), ("单价口径", "元/KG")])
        ws = wb[mod.NOTES_SHEET]
        self.assertEqual(ws["A1"].value, "导出时间")
        self.assertEqual(ws["B1"].value, "2026-09-21")
        self.assertEqual(ws["A2"].value, "单价口径")

    def test_header_only_when_no_data_rows(self):
        wb = self._write([])
        ws = wb[mod.DATA_SHEET]
        self.assertEqual(ws.max_row, 1)

    def test_auto_filter_spans_all_columns(self):
        wb = self._write([["A", "3011", "原料", 1.0, 2.5, None, "KG"]])
        self.assertEqual(wb[mod.DATA_SHEET].auto_filter.ref, "A1:G2")


# ======================================================================
# 命令级
# ======================================================================

class CommandTests(SimpleTestCase):

    # 显式给 output=None 表示「走默认 exports/ 路径」，
    # 与「没传 output」区分开，否则测不到默认命名
    _UNSET = object()

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.tmpdir = Path(self._tmp.name)

    def _run(self, out=None, **opts):
        out = out or io.StringIO()
        output = opts.pop("output", self._UNSET)
        master = opts.pop("master", master_df([master_row("A01005000057", "3011")]))
        stock = opts.pop("stock", stock_df([stock_row("A01005000057", "3011")]))
        price = opts.pop("price", price_df([price_row("A01005000057")]))

        def patch_query(name, value):
            """可调用对象当 side_effect（用于模拟 RFC 抛异常），否则当返回值。"""
            if callable(value):
                return patch.object(Command, name, side_effect=value)
            return patch.object(Command, name, return_value=value)

        kwargs = {"output": str(self.tmpdir / "out.xlsx")}
        if output is not self._UNSET:
            kwargs["output"] = output
        kwargs.update(opts)

        with patch(HEALTH_PATCH, return_value=HEALTH), \
                patch_query("_query_master", master), \
                patch_query("_query_stock", stock), \
                patch_query("_query_price", price):
            call_command("export_material_price_stock",
                         stdout=out, **kwargs)
        return out.getvalue()

    # ── 守卫 ──

    def test_rejects_zero_periods(self):
        with self.assertRaisesMessage(CommandError, "periods"):
            call_command("export_material_price_stock", periods=0)

    def test_rejects_negative_limit(self):
        with self.assertRaisesMessage(CommandError, "limit"):
            call_command("export_material_price_stock", limit=-1)

    def test_unhealthy_sap_aborts(self):
        with patch(HEALTH_PATCH, return_value={"status": "down", "error": "boom"}):
            with self.assertRaisesMessage(CommandError, "boom"):
                call_command("export_material_price_stock")

    def test_empty_stock_pull_aborts_without_writing(self):
        """配合「缺席记 0」，一次 RFC 故障会产出「库存全是 0」的假报告。"""
        target = self.tmpdir / "out.xlsx"
        with self.assertRaisesMessage(CommandError, "库存"):
            self._run(stock=pl.DataFrame(), output=str(target))
        self.assertFalse(target.exists())

    def test_empty_master_aborts(self):
        with self.assertRaisesMessage(CommandError, "物料"):
            self._run(master=master_df([]))

    def test_output_pointing_at_a_directory_is_rejected(self):
        with self.assertRaisesMessage(CommandError, "目录"):
            self._run(output=str(self.tmpdir))

    # ── 正常路径 ──

    def test_writes_file_with_expected_shape(self):
        target = self.tmpdir / "out.xlsx"
        stdout = self._run(output=str(target), periods=3)
        self.assertTrue(target.exists())
        self.assertIn("已写出", stdout)

        ws = load_workbook(target)[mod.DATA_SHEET]
        self.assertEqual(ws.max_row, 2)
        self.assertEqual(ws["A2"].value, "A01005000057")
        self.assertEqual(ws["B2"].value, "3011")
        self.assertEqual(ws["C2"].value, "原料X")
        self.assertEqual(ws["D2"].value, 10.0)
        self.assertEqual(ws["G2"].value, 12.5)   # 最后一列（当月）有价格
        self.assertEqual(ws["H2"].value, "KG")   # 单位列

    def test_dry_run_writes_nothing(self):
        target = self.tmpdir / "out.xlsx"
        stdout = self._run(output=str(target), dry_run=True)
        self.assertFalse(target.exists())
        self.assertIn("预览", stdout)

    def test_default_output_lands_in_export_dir(self):
        with patch.object(mod, "EXPORT_DIR", self.tmpdir / "exports"):
            stdout = self._run(output=None)
        written = list((self.tmpdir / "exports").glob("material_price_stock_*.xlsx"))
        self.assertEqual(len(written), 1)
        self.assertIn(str(written[0]), stdout)

    def test_output_path_overrides_timestamp_naming(self):
        target = self.tmpdir / "custom_name.xlsx"
        self._run(output=str(target))
        self.assertTrue(target.exists())

    def test_werks_filter_restricts_rows(self):
        """主数据侧也必须按工厂筛，否则非请求工厂的行会漏进结果。"""
        master = master_df([
            master_row("A1", "3011"), master_row("A1", "3020"),
        ])
        stock = stock_df([stock_row("A1", "3011"), stock_row("A1", "3020")])
        target = self.tmpdir / "out.xlsx"
        self._run(master=master, stock=stock, price=price_df([]),
                  werks="3011", output=str(target))
        ws = load_workbook(target)[mod.DATA_SHEET]
        self.assertEqual(ws.max_row, 2)
        self.assertEqual(ws["B2"].value, "3011")

    def test_only_with_data_drops_master_only_rows(self):
        master = master_df([
            master_row("A1", "3011"), master_row("A2", "3011"),
        ])
        target = self.tmpdir / "out.xlsx"
        self._run(master=master, stock=stock_df([stock_row("A1", "3011")]),
                  price=price_df([]), only_with_data=True, output=str(target))
        ws = load_workbook(target)[mod.DATA_SHEET]
        self.assertEqual(ws.max_row, 2)  # 只剩有库存的 A1
        self.assertEqual(ws["A2"].value, "A1")

    def test_exclude_deleted_drops_lvorm_rows(self):
        master = master_df([
            master_row("A1", "3011"), master_row("A2", "3011", lvorm="X"),
        ])
        target = self.tmpdir / "out.xlsx"
        stdout = self._run(master=master, stock=stock_df([stock_row("A1", "3011")]),
                           price=price_df([]), exclude_deleted=True,
                           output=str(target))
        ws = load_workbook(target)[mod.DATA_SHEET]
        self.assertEqual(ws.max_row, 2)
        self.assertEqual(ws["A2"].value, "A1")
        self.assertIn("删除标记", stdout)

    def test_ffill_fills_sparse_months(self):
        price = price_df([price_row("A01005000057", bdatj=2026, poper=9)])
        target = self.tmpdir / "out.xlsx"
        self._run(price=price, periods=3, ffill=True, output=str(target))
        ws = load_workbook(target)[mod.DATA_SHEET]
        # 窗口末月 2026-09；前两月被前值填充需要窗口内有更早数据，
        # 这里只有 9 月，所以只有末列有值 —— 断言不足以判定的情况不写死，
        # 改由 FillForwardTests 覆盖语义本身
        self.assertEqual(ws["G2"].value, 12.5)

    # ── 失败处理 ──

    def test_partial_price_failure_writes_file_then_errors(self):
        """等 40 秒不该什么都拿不到，但退出码必须是非零。"""
        target = self.tmpdir / "out.xlsx"

        def boom(matnr):
            raise SAPRfcError("ZRFC_GET_MBEWH", "connection reset")

        with self.assertRaises(CommandError):
            self._run(price=boom, output=str(target))
        self.assertTrue(target.exists())

    def test_allow_partial_suppresses_the_error(self):
        target = self.tmpdir / "out.xlsx"

        def boom(matnr):
            raise SAPRfcError("ZRFC_GET_MBEWH", "connection reset")

        self._run(price=boom, output=str(target), allow_partial=True)
        self.assertTrue(target.exists())

    def test_sap_business_error_counts_as_failure_not_empty(self):
        """只有 schema 声明过的空结果文案才算 empty，其余是真错误。"""
        def boom(matnr):
            raise SAPBusinessError(
                "ZRFC_GET_MBEWH", [SapMessage(type="E", text="授权失败")]
            )

        target = self.tmpdir / "out.xlsx"
        with self.assertRaises(CommandError):
            self._run(price=boom, output=str(target))
        wb = load_workbook(target)
        notes = [
            " ".join(str(v or "") for v in row)
            for row in wb[mod.NOTES_SHEET].iter_rows(values_only=True)
        ]
        self.assertTrue(any("价格查询失败" in n for n in notes))
        self.assertTrue(any("授权失败" in n for n in notes))

    def test_empty_price_result_is_not_a_failure(self):
        """SAP 无该物料价格（6~10% 命中率）属正常，不该报错退出。"""
        stdout = self._run(price=price_df([]))
        self.assertIn("SAP 无记录 1", stdout)

    # ── join 诊断 ──

    def test_join_diagnostics_are_reported(self):
        """数据侧多出主数据意味着 join 键坏了，必须喊出来。"""
        master = master_df([master_row("A1", "3011")])
        stock = stock_df([stock_row("A1", "3011"), stock_row("A2", "3011")])
        stdout = self._run(master=master, stock=stock)
        self.assertIn("join", stdout)

    def test_reports_unmatched_bwkey(self):
        master = master_df([master_row("A01005000057", "3011")])
        price = price_df([price_row("A01005000057", bwkey="9999")])
        target = self.tmpdir / "out.xlsx"
        stdout = self._run(master=master, price=price, output=str(target))
        self.assertIn("9999", stdout)

    def test_whitespace_padded_werks_still_joins(self):
        """主数据 WERKS 带尾空格 + 价格 BWKEY 没有 → 归一化后仍能对上。

        不归一化的话价格列会整列变空，且没有任何报错。
        """
        master = master_df([master_row("A01005000057", "3011 ")])
        target = self.tmpdir / "out.xlsx"
        self._run(master=master, periods=3, output=str(target))
        ws = load_workbook(target)[mod.DATA_SHEET]
        self.assertEqual(ws["B2"].value, "3011")
        self.assertEqual(ws["G2"].value, 12.5)


# ======================================================================
# _query_master 的段合并与守卫
# ======================================================================

class QueryMasterTests(SimpleTestCase):

    def test_empty_pattern_does_not_break_concat(self):
        """空结果是零列帧，与有列帧 concat 会抛 ShapeError。"""
        frames = {
            "A01*": master_df([master_row("A1", "3011")]),
            "1002*": pl.DataFrame(),  # 零列空帧
        }
        with patch.object(Command, "_query_master_pattern",
                          side_effect=lambda p, w: frames[p]):
            out = Command()._query_master(("A01*", "1002*"), [])
        self.assertEqual(out.height, 1)

    def test_one_pattern_failing_aborts(self):
        """缺一段会静默丢掉整段物料，必须中止。"""
        def flaky(pattern, werks):
            if pattern == "1002*":
                raise RuntimeError("timeout")
            return master_df([master_row("A1", "3011")])

        with patch.object(Command, "_query_master_pattern", side_effect=flaky):
            with self.assertRaisesMessage(CommandError, "1002"):
                Command()._query_master(("A01*", "1002*"), [])

    def test_all_patterns_empty_aborts(self):
        with patch.object(Command, "_query_master_pattern",
                          side_effect=lambda p, w: pl.DataFrame()):
            with self.assertRaisesMessage(CommandError, "空结果"):
                Command()._query_master(("A01*",), [])
