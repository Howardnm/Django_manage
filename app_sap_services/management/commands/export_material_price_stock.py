"""
Django 管理命令: 实时从 SAP 导出原材料「库存 + 逐月价格」Excel 报表。

用法:
    python manage.py export_material_price_stock                      # 近 12 个月，全量
    python manage.py export_material_price_stock --periods 24         # 近 24 个月
    python manage.py export_material_price_stock --matnr "A01*"       # 只导某料号段
    python manage.py export_material_price_stock --matnr A01005000057 # 只导单个物料
    python manage.py export_material_price_stock --werks 3011,3020    # 限定工厂
    python manage.py export_material_price_stock --limit 20 --verbose # 小样本试跑
    python manage.py export_material_price_stock --dry-run            # 只统计，不写文件
    python manage.py export_material_price_stock --output D:/out.xlsx # 指定输出路径

输出列（每行一个 物料×工厂）:

    原材料SAP编码 | 工厂 | 原材料名称 | 原材料实时库存 | 2025-10 | 2025-11 | ... | 基本计量单位

取数（全部实时，不读本地库）:
    主数据   ZRFC_MATERIAL_MESN   按 SYNC_PATTERNS 逐段拉，产出物料清单/名称/单位
    库存     ZRFC_GET_MAT_STOCK   一次全量（实测约 1 秒），按 物料×工厂 汇总 CLABS
    价格     ZRFC_GET_MBEWH       逐物料遍历（2654 个物料约 25~50 秒）

    价格必须逐物料拉：ZRFC_GET_MBEWH 不传物料筛选时实测 556 万行、近 10 分钟。
    这条策略与 sync_material_prices 一致，那边有更详细的取舍说明。

库存口径:
    只取非限制库存 CLABS，按 (物料, 工厂) 汇总，不含安全库存、不拆库位/批次。
    ZRFC_GET_MAT_STOCK 的通配查询不返回零库存物料，因此「结果里没有」= 库存为 0，
    与 sync_material_stock 的 zero_fill 判断一致。

价格口径:
    UNIT_PRICE = VERPR / PEINH（实测 STPRS/PVPRS 恒为 0，只有 VERPR 有值）。
    单位是「元/基本计量单位」而非一定元/kg —— 所以末尾带一列 MEINS。
    空白月份表示 MBEWH 在该期间没有估值变动记录，**不等于没有价格**
    （MBEWH 只在估值变化时记行，数据天然稀疏）。
    清洗逻辑与写库路径共用 transforms.transform_prices，但传
    price_decimals=None 保留原商 —— 写库路径先 round(2) 再判 > 0，PEINH=10000
    时会把单价 < 0.005 的真实价格舍成 0 丢弃（见该函数 docstring）。

写入语义:
    只读导出，不写数据库、不建 Plant 记录。未知工厂代码原样输出。
    只要 SAP 拉取成功就一定会产出文件；逐物料价格失败会记入说明页并在最后
    以非零退出码收尾（--allow-partial 可抑制），避免「等 40 秒却什么都没拿到」。
"""

from __future__ import annotations

from collections import Counter, defaultdict
from pathlib import Path

import polars as pl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.utils import timezone

from app_sap_services import (
    SAPBusinessError,
    SAPError,
    sap,
    sap_health_check,
)
from app_sap_services.definitions.material import MaterialQuery
from app_sap_services.definitions.price import MaterialPriceQuery
from app_sap_services.definitions.stock import MaterialStockQuery
from app_sap_services.transforms import (
    apply_period_filter,
    blank,
    clean_stock,
    current_ym,
    month_window,
    normalize_keys,
    period_ym,
    transform_prices,
    ym_label,
)


# 与 sync_raw_materials 保持一致：每段单独下发到 SAP 服务端筛选。
# 不要图省事合并成单次 "A0*" 宽拉 —— 1001*/1002* 段将永远拿不到数据。
SYNC_PATTERNS = ("A01*", "A03*", "1001*", "1002*")

DEFAULT_PERIODS = 12
PROGRESS_EVERY = 200

EXPORT_DIR = Path(settings.BASE_DIR) / "exports"
DATA_SHEET = "价格库存"
NOTES_SHEET = "说明"

HEAD_CODE = "原材料SAP编码"
HEAD_PLANT = "工厂"
HEAD_NAME = "原材料名称"
HEAD_STOCK = "原材料实时库存"
HEAD_UNIT = "基本计量单位"

FMT_TEXT = "@"
FMT_STOCK = "0.000"
FMT_PRICE = "0.0000"


# ======================================================================
# 纯函数：行集合 / 透视 / 写出（可脱离 SAP 与 DB 单测）
# ======================================================================

def build_rows(master_pairs, stock_by_pair, price_index, only_with_data=False):
    """确定导出的 (物料, 工厂) 行集合，按 (物料, 工厂) 升序。

    取 master ∪ 库存 ∪ 价格 三者的并集：任何一侧有数据就不该被丢掉。
    正常情况下并集会等于 master —— 数据侧多出 master 的部分几乎一定是
    join key 出问题（补空格、BWKEY≠WERKS），调用方应把差值报出来当报警。

    Args:
        master_pairs: 主数据里的 (MATNR, WERKS) 可迭代对象。
        stock_by_pair / price_index: 以 (matnr, plant) 为键的映射。
        only_with_data: True 时丢弃既无库存、窗口内又无价格的行。
    """
    pairs = set(master_pairs)
    data_pairs = set(stock_by_pair) | set(price_index)
    pairs |= data_pairs
    if only_with_data:
        pairs &= data_pairs
    return sorted(pairs, key=lambda p: (p[0], p[1]))


def fill_forward(by_month, months):
    """在窗口内向后续月份填充最近的已知价格（--ffill）。

    只在窗口内有数据的月份之间填充；窗口开始前的数据不参与，
    也不向前填充（那会凭空造出未来价格）。
    """
    out = {}
    last = None
    for month in months:
        if month in by_month:
            last = by_month[month]
        if last is not None:
            out[month] = last
    return out


def build_data_rows(pairs, months, names, units, stock_by_pair, price_index,
                    forward_fill=False):
    """把行集合 + 各维度映射摊平成可直接写入 Excel 的二维列表。

    缺失月份写 None（留空），**不是 0** —— 空白表示「该期间无估值变动记录」，
    与「价格为 0」是两回事。库存缺席则写 0.0，理由见模块 docstring。
    """
    rows = []
    for pair in pairs:
        by_month = price_index.get(pair, {})
        if forward_fill:
            by_month = fill_forward(by_month, months)
        matnr = pair[0]
        rows.append(
            [matnr, pair[1], names.get(matnr, ""), stock_by_pair.get(pair, 0.0)]
            + [by_month.get(m) for m in months]
            + [units.get(matnr, "")]
        )
    return rows


def _safe_append(ws, values):
    """append 一行，并把以 '=' 开头的字符串钉成文本。

    SAP 的物料描述是人填的，以 '=' 开头会被 openpyxl 当成公式
    （data_type='f'，Excel 里显示 #NAME?），既是渲染问题也是注入面。
    """
    ws.append(values)
    row_idx = ws.max_row
    for col_idx, value in enumerate(values, start=1):
        if isinstance(value, str) and value.startswith("="):
            ws.cell(row=row_idx, column=col_idx).data_type = "s"


def write_workbook(dest, month_labels, data_rows, notes):
    """写出报表。

    Args:
        dest: 文件路径或可写流（测试传 io.BytesIO 即可，无需落盘）。
        month_labels: 月份列表头，如 ['2025-10', ...]，升序。
        data_rows: build_data_rows 的产物，每行
                   [编码, 工厂, 名称, 库存, 各月价格..., 单位]。
        notes: [(标签, 内容), ...]，写进「说明」页。
    """
    headers = (
        [HEAD_CODE, HEAD_PLANT, HEAD_NAME, HEAD_STOCK]
        + list(month_labels)
        + [HEAD_UNIT]
    )
    n_cols = len(headers)
    last_col = get_column_letter(n_cols)

    wb = Workbook()
    ws = wb.active
    ws.title = DATA_SHEET
    _safe_append(ws, headers)

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="D9D9D9")
    header_align = Alignment(horizontal="center", vertical="center")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    for values in data_rows:
        _safe_append(ws, values)

    # 月份列多，冻结到库存列之后 —— 冻 A2 的话横滚就看不见编码了
    ws.freeze_panes = "E2"

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=n_cols):
        for cell in row:
            if cell.column <= 2:
                cell.number_format = FMT_TEXT
            elif cell.column == 4:
                cell.number_format = FMT_STOCK
            elif cell.column < n_cols:
                cell.number_format = FMT_PRICE

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 8
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 16
    for i in range(5, n_cols):
        ws.column_dimensions[get_column_letter(i)].width = 12
    ws.column_dimensions[last_col].width = 10

    # 数据行可能为 0，此时只框住表头
    ws.auto_filter.ref = f"A1:{last_col}{max(ws.max_row, 1)}"

    notes_ws = wb.create_sheet(NOTES_SHEET)
    notes_ws.column_dimensions["A"].width = 24
    notes_ws.column_dimensions["B"].width = 96
    for label, value in notes:
        notes_ws.append([label, value])
        notes_ws.cell(row=notes_ws.max_row, column=1).font = Font(bold=True)

    wb.save(dest)


def _split_csv(value):
    """'3011, 3020' → ['3011', '3020']；空 → []"""
    if not value:
        return []
    return [item.strip() for item in value.split(",") if item.strip()]


def _ordered_unique(values):
    return list(dict.fromkeys(values))


# ======================================================================
# 命令
# ======================================================================

class Command(BaseCommand):
    help = "实时从 SAP 导出原材料库存 + 逐月价格 Excel 报表（不写数据库）"

    def add_arguments(self, parser):
        parser.add_argument(
            "--periods", type=int, default=DEFAULT_PERIODS,
            help=f"月价格列数，近 N 个月（含当月，默认 {DEFAULT_PERIODS}）",
        )
        parser.add_argument(
            "--werks", "--bwkey", dest="werks", type=str, default=None,
            help="工厂筛选，逗号分隔（如 '3011,3020'）。多工厂在客户端过滤",
        )
        parser.add_argument(
            "--matnr", type=str, default=None,
            help="物料匹配模式，支持 * 通配（如 'A01*'）。默认 SYNC_PATTERNS 全段",
        )
        parser.add_argument(
            "--output", type=str, default=None,
            help="输出文件的完整路径（给了就不加时间戳、不落 exports/ 目录）",
        )
        parser.add_argument(
            "--limit", type=int, default=0,
            help="只处理排序后的前 N 个物料（抽样试跑，结果不完整）",
        )
        parser.add_argument(
            "--dry-run", action="store_true",
            help="只查询并统计，不写文件",
        )
        parser.add_argument(
            "--only-with-data", action="store_true", dest="only_with_data",
            help="丢弃既无库存、窗口内又无价格的 (物料,工厂) 行",
        )
        parser.add_argument(
            "--exclude-deleted", action="store_true", dest="exclude_deleted",
            help="丢弃 SAP 删除标记 (LVORM) 非空的物料-工厂行（默认保留并计数）",
        )
        parser.add_argument(
            "--ffill", action="store_true",
            help="窗口内稀疏月份按前值填充（默认留空）",
        )
        parser.add_argument(
            "--allow-partial", action="store_true", dest="allow_partial",
            help="逐物料价格查询失败时仍返回 0 退出码（默认写文件后报错退出）",
        )
        parser.add_argument(
            "--verbose", action="store_true",
            help="逐物料打印进度明细",
        )

    # ------------------------------------------------------------------
    # 入口
    # ------------------------------------------------------------------

    def handle(self, *args, **options):
        periods = options["periods"]
        limit = options["limit"]
        dry_run = options["dry_run"]
        only_with_data = options["only_with_data"]
        exclude_deleted = options["exclude_deleted"]
        forward_fill = options["ffill"]
        allow_partial = options["allow_partial"]
        verbose = options["verbose"]
        output = options["output"]
        matnr = options["matnr"]
        werks = _split_csv(options["werks"])
        single_plant = werks[0] if len(werks) == 1 else None

        if periods < 1:
            raise CommandError("--periods 必须 >= 1（月列数）")
        if limit < 0:
            raise CommandError("--limit 不能为负数")

        patterns = (matnr,) if matnr else SYNC_PATTERNS

        self.stdout.write(
            self.style.MIGRATE_HEADING(
                f"\n=== SAP 原材料价格/库存导出 ===\n"
                f"    料号段: {', '.join(patterns)}\n"
                f"    工厂:   {', '.join(werks) if werks else '全部'}\n"
                f"    窗口:   近 {periods} 个月"
            )
        )
        if dry_run:
            self.stdout.write(
                self.style.WARNING("   [DRY-RUN] 只统计，不写文件")
            )

        # ── 1. 健康检查 ──
        health = sap_health_check()
        if health.get("status") != "healthy":
            raise CommandError(f"SAP 连接失败: {health.get('error', '未知错误')}")
        self.stdout.write(
            f"   [OK] SAP 连接正常 "
            f"(ashost={health.get('ashost')}, client={health.get('client')})"
        )

        # ── 2. 月份窗口 ──
        today = timezone.localdate()
        end_ym = current_ym(today)
        months = month_window(end_ym, periods)
        ym_set = set(months)
        self.stdout.write(
            f"   月份窗口: {ym_label(months[0])} ~ {ym_label(months[-1])}"
        )

        # ── 3. 主数据 ──
        master = self._query_master(patterns, werks)
        master = normalize_keys(master, ["MATNR", "WERKS", "MAKTX", "MEINS"])
        master = master.filter(pl.col("MATNR") != "")
        master = master.filter(pl.col("WERKS") != "")
        if werks:
            # 单工厂时服务端已筛过，这里是幂等的兜底；多工厂只能客户端筛
            # （RfcQuery 每个 range 只持一行）。不筛的话，非请求工厂的行会
            # 从主数据侧漏进结果。
            master = master.filter(pl.col("WERKS").is_in(werks))
        if blank(master):
            raise CommandError("主数据清洗后无有效行（物料号/工厂为空），已中止")

        deleted_rows = master.filter(pl.col("LVORM") != "").height
        master = master.filter(pl.col("LVORM") == "") if exclude_deleted else master

        names = self._first_non_empty(master, "MAKTX")
        units = self._first_non_empty(master, "MEINS")
        material_ids = sorted(_ordered_unique(master["MATNR"].to_list()))
        if limit and limit > 0:
            material_ids = material_ids[:limit]
        # 行集合从**已应用 --exclude-deleted** 的主数据派生，再按最终物料清单收窄。
        # 换成过滤前的集合会让被排除的删除标记行重新漏回结果。
        codes = set(material_ids)
        master_pairs = {
            pair
            for pair in master.select(["MATNR", "WERKS"]).unique().iter_rows()
            if pair[0] in codes
        }

        self.stdout.write(
            f"   主数据: {master.height} 行 → 物料 {len(material_ids)} 个, "
            f"物料×工厂 {len(master_pairs)} 对"
            + (f", 删除标记 {deleted_rows} 行" if deleted_rows else "")
        )
        if not material_ids:
            raise CommandError("主数据里没有可导出的物料，已中止")

        # ── 4. 库存 ──
        stock_warnings = Counter()
        stock_df = self._query_stock(werks)
        if blank(stock_df):
            # 配合「缺席记为 0」的策略，一次 RFC 故障会产出「所有物料库存都是 0」
            # 的假报告，比同步场景更危险，所以这里必须中止
            raise CommandError(
                "库存全量查询返回空结果，已中止（疑似 RFC 故障，未写文件）"
            )
        stock_df = clean_stock(stock_df, stock_warnings)
        stock_df = normalize_keys(stock_df, ["MATNR", "WERKS"])
        stock_df = stock_df.filter(pl.col("MATNR").is_in(material_ids))
        if werks:
            stock_df = stock_df.filter(pl.col("WERKS").is_in(werks))
        stock_by_pair = {
            (row["MATNR"], row["WERKS"]): float(row["CLABS"] or 0)
            for row in stock_df.group_by(["MATNR", "WERKS"])
            .agg(pl.col("CLABS").sum())
            .iter_rows(named=True)
        }
        stock_names = self._first_non_empty(stock_df, "MAKTX")
        names = {**stock_names, **names}  # 主数据的名称优先
        self.stdout.write(
            f"   库存: {stock_df.height} 行 → {len(stock_by_pair)} 对物料×工厂"
        )

        # ── 5. 价格（逐物料遍历）──
        price_index, price_stats, failed_codes = self._collect_prices(
            material_ids, months, ym_set, single_plant, verbose
        )

        # ── 6. 行集合 ──
        pairs = build_rows(
            master_pairs, stock_by_pair, price_index,
            only_with_data=only_with_data,
        )
        # 以 master_pairs（已按 --limit / --exclude-deleted 收窄）为基准比较，
        # 否则这两个开关会让「数据侧多出主数据」凭空出现假报警
        data_pairs = set(stock_by_pair) | set(price_index)
        extra_pairs = data_pairs - master_pairs
        missing_pairs = master_pairs - data_pairs
        unmatched_bwkey = sorted(
            {pair[1] for pair in price_index} - {pair[1] for pair in master_pairs}
        )

        # ── 7. 写出 ──
        month_labels = [ym_label(m) for m in months]
        data_rows = build_data_rows(
            pairs, months, names, units, stock_by_pair, price_index,
            forward_fill=forward_fill,
        )
        notes = self._build_notes(
            health=health, months=months, periods=periods, patterns=patterns,
            werks=werks, today=today, row_count=len(pairs),
            month_count=len(months), material_count=len(
                {p[0] for p in pairs}),
            plant_count=len({p[1] for p in pairs}),
            price_stats=price_stats, stock_warnings=stock_warnings,
            failed_codes=failed_codes, extra_pairs=extra_pairs,
            missing_pairs=missing_pairs, unmatched_bwkey=unmatched_bwkey,
            deleted_rows=deleted_rows, exclude_deleted=exclude_deleted,
            only_with_data=only_with_data, forward_fill=forward_fill,
            limit=limit,
        )

        output_path = None
        if not dry_run:
            output_path = self._resolve_output(output)
            write_workbook(output_path, month_labels, data_rows, notes)
            self.stdout.write(
                self.style.SUCCESS(f"\n[OK] 已写出: {output_path}")
            )

        self._report(
            rows=len(pairs), months=len(months), data_rows=data_rows,
            price_stats=price_stats, failed_codes=failed_codes,
            stock_by_pair=stock_by_pair, price_index=price_index,
            extra_pairs=extra_pairs, missing_pairs=missing_pairs,
            unmatched_bwkey=unmatched_bwkey, dry_run=dry_run,
        )

        if failed_codes and not allow_partial:
            # 文件已经写出（人力等待 40 秒不该什么都拿不到），但退出码必须是
            # 非零，否则 Task Scheduler 会把一次残缺导出当成成功
            raise CommandError(
                f"有 {len(failed_codes)} 个物料的价格查询失败，"
                f"结果不完整（详见上方日志与说明页）。"
                f"如需忽略请加 --allow-partial"
            )

    # ------------------------------------------------------------------
    # SAP 取数（三个可 patch 的 I/O 缝合点）
    # ------------------------------------------------------------------

    def _query_master_pattern(self, pattern, werks) -> pl.DataFrame:
        """单个料号段的主数据查询。"""
        query = sap.rfc(MaterialQuery).filter(mat_range__cp=pattern)
        if len(werks) == 1:
            # RfcQuery.filter 每个 range 只能持一行（内部是 dict.update），
            # 多工厂表达不了，只能在客户端过滤
            query = query.filter(wek_range__eq=werks[0])
        return query.collect()

    def _query_master(self, patterns, werks) -> pl.DataFrame:
        """逐段拉主数据并合并。

        任一段失败即中止：主数据缺一段会静默丢掉整个料号段，
        报出来的表看起来完全正常，只是少了几千行。
        """
        chunks = []
        failures = []
        for pattern in patterns:
            try:
                part = self._query_master_pattern(pattern, werks)
            except Exception as e:
                failures.append((pattern, str(e)))
                continue
            if not part.is_empty():
                # 空结果是零列帧，与有列帧 concat 会抛 ShapeError
                chunks.append(part)
            self.stdout.write(f"   [{pattern}] SAP 返回: {part.height} 条")

        if failures:
            detail = "; ".join(f"{p}: {e}" for p, e in failures)
            raise CommandError(
                f"有 {len(failures)}/{len(patterns)} 个料号段查询失败，已中止"
                f"（缺段会静默丢掉整段物料）: {detail}"
            )
        if not chunks:
            raise CommandError("所有料号段都返回空结果，已中止（疑似 RFC 故障）")
        return pl.concat(chunks)

    def _query_stock(self, werks) -> pl.DataFrame:
        """库存查询。单工厂下推服务端，多工厂走全量再客户端过滤。"""
        query = sap.rfc(MaterialStockQuery)
        if len(werks) == 1:
            query = query.filter(wek_range__eq=werks[0])
        else:
            query = query.filter(mat_range__cp="*")
        df = query.collect()
        self.stdout.write(f"   SAP 库存返回: {df.height} 条")
        return df

    def _query_price(self, matnr) -> pl.DataFrame:
        """单物料价格查询（该物料全工厂全历史）。"""
        if "*" in matnr:
            filters = {"s_matnr__cp": matnr}
        else:
            filters = {"s_matnr__eq": matnr}
        return sap.rfc(MaterialPriceQuery).filter(**filters).collect()

    # ------------------------------------------------------------------
    # 价格收集
    # ------------------------------------------------------------------

    def _collect_prices(self, material_ids, months, ym_set, single_plant, verbose):
        """逐物料拉价格，摊平成 {(物料, 工厂): {YYYYMM: 单价}}。"""
        price_index = defaultdict(dict)
        stats = Counter()
        failed_codes = []
        lo, hi = months[0], months[-1]
        window_end_ym = months[-1]

        for idx, code in enumerate(material_ids, 1):
            try:
                df = self._query_price(code)
            except SAPBusinessError as e:
                # 声明过的「未查询到相关数据」不会走到这里（schema 已识别为空结果）
                stats["failed"] += 1
                failed_codes.append((code, str(e)))
                self.stderr.write(self.style.ERROR(f"   [{code}] SAP 业务错误: {e}"))
                continue
            except SAPError as e:
                stats["failed"] += 1
                failed_codes.append((code, str(e)))
                self.stderr.write(self.style.ERROR(f"   [{code}] SAP 调用失败: {e}"))
                continue

            if blank(df):
                stats["empty"] += 1
                continue

            # 以下统计必须在清洗之前做，清洗会丢掉这些行
            stats["vprsv_s"] += df.filter(pl.col("VPRSV") == "S").height
            stats["odd_period"] += df.filter(
                pl.col("POPER").is_null()
                | (pl.col("POPER") < 1)
                | (pl.col("POPER") > 12)
            ).height

            df = normalize_keys(df, ["MATNR", "BWKEY"])
            # 窗口末月即当前月；区间过滤已排除窗口外，未来期间护栏在此
            # 只负责把「晚于当前月」的记录计进 future 计数
            df = apply_period_filter(
                df, (lo, hi),
                current_ym=window_end_ym,
                include_future=False,
                warnings=stats,
            )
            df = transform_prices(df, stats, price_decimals=None)

            if df.is_empty():
                stats["no_price"] += 1
                continue

            for row in df.iter_rows(named=True):
                bwkey = row["BWKEY"]
                if single_plant and bwkey != single_plant:
                    continue
                ym = period_ym(row["BDATJ"], row["POPER"])
                if ym in ym_set:
                    # 保留原商避免 <0.005 被吞（见 transform_prices），但存成
                    # 6 位小数以免浮点噪声（PEINH=10000 时原商会算出
                    # 7.887719000000001 这种值），Excel 显示 4 位
                    price_index[(row["MATNR"], bwkey)][ym] = round(
                        float(row["UNIT_PRICE"]), 6
                    )
            stats["with_price"] += 1

            if verbose:
                self.stdout.write(
                    f"   [{code}] {df.height} 条 → 累计 "
                    f"{len(price_index)} 对物料×工厂"
                )
            elif idx % PROGRESS_EVERY == 0:
                self.stdout.write(
                    f"   进度: {idx}/{len(material_ids)} "
                    f"（有价格 {stats['with_price']} 个, "
                    f"SAP 无记录 {stats['empty']} 个, "
                    f"失败 {stats['failed']} 个）"
                )

        return price_index, stats, failed_codes

    # ------------------------------------------------------------------
    # 辅助
    # ------------------------------------------------------------------

    @staticmethod
    def _first_non_empty(df, column) -> dict:
        """{物料: 该列第一个非空值}，按 (MATNR, WERKS) 排序保证可复现。

        同一物料的 MAKTX/MEINS 在 MAKT/MARA 里与工厂无关，取哪个都一样；
        排序是为了让「取第一个」这个动作在不同机器上结果一致。
        """
        if blank(df) or column not in df.columns:
            return {}
        ordered = df.sort(["MATNR", "WERKS"]) if "WERKS" in df.columns else df
        result = {}
        for row in ordered.iter_rows(named=True):
            matnr = row["MATNR"]
            if result.get(matnr):
                continue
            value = (row.get(column) or "").strip()
            if value:
                result[matnr] = value
        return result

    @staticmethod
    def _resolve_output(output) -> Path:
        if output:
            path = Path(output)
            if path.is_dir():
                raise CommandError(
                    f"--output 指向的是一个目录，请给出完整文件路径: {path}"
                )
            path.parent.mkdir(parents=True, exist_ok=True)
            return path
        EXPORT_DIR.mkdir(parents=True, exist_ok=True)
        stamp = timezone.localtime().strftime("%Y%m%d_%H%M%S")
        return EXPORT_DIR / f"material_price_stock_{stamp}.xlsx"

    def _build_notes(self, **kw):
        """说明页内容 —— 报表脱离上下文后，这些口径没法从数字里读出来。"""
        health = kw["health"]
        months = kw["months"]
        stats = kw["price_stats"]
        stock_warnings = kw["stock_warnings"]

        notes = [
            ("导出时间", timezone.localtime().strftime("%Y-%m-%d %H:%M:%S")),
            ("数据来源", "实时调用 SAP RFC（不读取本地数据库）"),
            ("SAP 系统", f"ashost={health.get('ashost')} "
                         f"client={health.get('client')}"),
            ("RFC 函数",
             "ZRFC_MATERIAL_MESN（主数据）/ ZRFC_GET_MAT_STOCK（库存）"
             "/ ZRFC_GET_MBEWH（价格）"),
            ("料号段", ", ".join(kw["patterns"])),
            ("工厂筛选", ", ".join(kw["werks"]) if kw["werks"] else "全部"),
            ("月份窗口",
             f"{ym_label(months[0])} ~ {ym_label(months[-1])}"
             f"（近 {kw['periods']} 个月，含当月 {kw['today']}）"),
            ("行口径", f"{kw['row_count']} 行 = 物料 {kw['material_count']} 个"
                       f" × 工厂 {kw['plant_count']} 个（每行一对）"),
            ("列口径", f"4 个基础列 + {kw['month_count']} 个月份列 + 1 个单位列"),
            ("单价口径",
             "UNIT_PRICE = VERPR / PEINH，单位是「元/基本计量单位」"
             "（见末列，可能是 KG/L/PC，不一定是 kg）。"
             "显示 4 位小数；判断有效性用的是未舍入的原商"),
            ("库存口径",
             "只取非限制库存 CLABS，按 物料×工厂 汇总（不含安全库存、"
             "不拆库位/批次）。SAP 通配查询不返回零库存物料，"
             "因此表中缺席记 0"),
            ("空白月份的含义",
             "MBEWH 只在估值发生变化时记行，空白 = 该期间没有估值变动记录，"
             "**不等于没有价格**"),
        ]

        if kw["forward_fill"]:
            notes.append(
                ("前值填充", "已启用 --ffill：窗口内空白月份按前一有效价格填充")
            )
        if kw["only_with_data"]:
            notes.append(
                ("行筛选", "已启用 --only-with-data：既无库存、窗口内又无价格的行已丢弃")
            )
        if kw["limit"]:
            notes.insert(
                0, ("⚠ 抽样导出",
                    f"使用了 --limit {kw['limit']}，只处理了前 {kw['limit']} 个物料，"
                    f"**结果不完整**，不可用于全量分析")
            )

        notes.extend([
            ("价格查询",
             f"有价格 {stats['with_price']} 个物料, "
             f"SAP 无记录 {stats['empty']} 个（实测 6~10%，属正常）, "
             f"有记录但无有效价格 {stats['no_price']} 个, "
             f"失败 {stats['failed']} 个"),
            ("被过滤的记录",
             f"非 CNY 币种 {stats['currency']} 条, "
             f"重复期间（按 KALNR 取最大）{stats['dedupe']} 条, "
             f"未来期间 {stats['future']} 条, "
             f"POPER 非 1-12（调整期间等）{stats['odd_period']} 条"),
            ("库存清洗",
             f"空工厂哨兵行 {stock_warnings['empty_plant']} 条已丢弃"),
            ("VPRSV='S' 行数",
             f"{stats['vprsv_s']}。已核对：STPRS/PVPRS 在实测抽样中恒为 0"
             f"（含 VPRSV='S' 的行），故标准价控制物料同样只有 VERPR 可取值，"
             f"口径不受影响。此计数保留作 SAP 配置变更的哨兵"),
            ("删除标记",
             f"LVORM 非空 {kw['deleted_rows']} 行，"
             + ("已按 --exclude-deleted 排除" if kw["exclude_deleted"]
                else "已保留（未加 --exclude-deleted）")),
            ("join 校验",
             f"数据侧多出主数据的 (物料,工厂) {len(kw['extra_pairs'])} 对"
             f"（不为 0 说明 join 键有问题）; "
             f"主数据中窗口内无任何数据的 {len(kw['missing_pairs'])} 对"),
        ])

        if kw["unmatched_bwkey"]:
            notes.append(
                ("未匹配的评估范围",
                 f"{', '.join(kw['unmatched_bwkey'][:10])}"
                 f"{' …' if len(kw['unmatched_bwkey']) > 10 else ''}"
                 f"（价格有数据但主数据里没有对应工厂）")
            )
        if kw["failed_codes"]:
            sample = "; ".join(f"{c}: {e}" for c, e in kw["failed_codes"][:10])
            notes.append(
                ("⚠ 价格查询失败",
                 f"{len(kw['failed_codes'])} 个物料，这些物料的价格列为空。"
                 f"示例: {sample}")
            )

        notes.extend([
            ("假设 1", "会计年度 BDATJ 与自然年对齐（窗口按 YYYYMM 比较）"),
            ("假设 2", "评估范围 BWKEY 视为工厂，与 WERKS 同一编码空间"),
            ("使用提示", "本表为只读快照，导入回系统请勿直接使用"),
        ])
        return notes

    def _report(self, rows, months, data_rows, price_stats, failed_codes,
                stock_by_pair, price_index, extra_pairs, missing_pairs,
                unmatched_bwkey, dry_run):
        prefix = "将写入" if dry_run else "已写入"
        self.stdout.write("")
        self.stdout.write(
            self.style.SUCCESS(
                f"[OK] 导出{'预览' if dry_run else '完成'}！"
                f"{prefix} {rows} 行 × {months} 个月份列"
            )
        )
        priced = sum(1 for r in data_rows if any(v is not None for v in r[4:4 + months]))
        self.stdout.write(
            f"     有价格的行: {priced}/{rows}，"
            f"有库存的行: {sum(1 for r in data_rows if r[3])}/{rows}"
        )
        self.stdout.write(
            f"     物料 {len({p[0] for p in set(stock_by_pair) | set(price_index)})} 个, "
            f"物料×工厂 {len(set(stock_by_pair) | set(price_index))} 对"
        )
        self.stdout.write(
            f"     价格查询: 有价格 {price_stats['with_price']} / "
            f"SAP 无记录 {price_stats['empty']} / "
            f"无有效价格 {price_stats['no_price']} / "
            f"失败 {price_stats['failed']}"
        )
        if price_stats["vprsv_s"]:
            self.stdout.write(
                self.style.WARNING(
                    f"     [注意] VPRSV='S' 的价格行 {price_stats['vprsv_s']} 条；"
                    f"已核对 STPRS 恒为 0，VERPR 口径不受影响（计数保留作哨兵）"
                )
            )
        if extra_pairs:
            self.stdout.write(
                self.style.WARNING(
                    f"     [警告] 数据侧多出主数据 {len(extra_pairs)} 对 (物料,工厂)，"
                    f"疑似 join 键（空格/BWKEY≠WERKS）问题，请检查"
                )
            )
        if unmatched_bwkey:
            self.stdout.write(
                self.style.WARNING(
                    f"     [警告] 未匹配到主数据的评估范围: "
                    f"{', '.join(unmatched_bwkey[:10])}"
                )
            )
        if failed_codes:
            self.stdout.write(self.style.ERROR("     失败示例（最多 10 个）:"))
            for code, err in failed_codes[:10]:
                self.stdout.write(self.style.ERROR(f"       {code}: {err}"))
