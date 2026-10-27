"""按项目编号分步同步 init/项目导入.xlsx，并把结果写回「同步情况」。

成功时写「已同步」；关联不上的项追加为「已同步-异常1；异常2」。
这些关联失败不回滚本行。项目建不起来（没有负责人、状态无法识别等）仍写异常并回滚。

    python manage.py import_projects
    python manage.py import_projects --limit 5
    python manage.py import_projects --step project --step repository
    python manage.py import_projects --dry-run
    python manage.py import_projects --export
    python manage.py import_projects --export init/项目导出.xlsx
"""
import os
import secrets
import string
from datetime import date, datetime

from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.db.models.signals import post_save
from django.utils import timezone

from app_material.models import MaterialLibrary, MaterialType
from app_project.models import BusinessSegment, Project, ProjectConfig, ProjectNode, ProjectStage
from app_repository.models import Customer, GradeFactor, ProjectRepository

User = get_user_model()

STEP_ORDER = ('project', 'repository', 'material', 'progress')
STEP_RANK = {'project': 1, 'repository': 2, 'material': 3, 'progress': 4}
RANK_MARK = {1: '项目已同步', 2: '档案已同步', 3: '材料已同步', 4: '已同步'}
SYNCED = '已同步'
GRADE_PREFIX = '请编辑牌号-'
GRADE_ALPHABET = string.ascii_letters + string.digits

HEADERS = (
    '同步情况',
    '项目编号',
    '项目名称',
    '项目等级',
    '*成品SAP码',
    '材料牌号',
    '*材料类型',
    '*业务板块',
    '*研发负责人',
    '立项时间',
    '*项目已完成状态',
    '项目状态',
    '项目描述：异常原因',
    '项目描述：备注',
    '项目描述：不合格原因',
    '项目描述：总出货数量（吨）',
    '*客户编码',
    '*业务负责人',
    '制件名称(零件名或总成)',
    '制件用途（用在哪个产品）',
    '制件颜色（项目档案）',
    '研发高可落地性项目',
    '业务高可落地性项目',
)

# 「*项目已完成状态」+「项目状态」→ 要标成已完成的阶段。下一节点另设为进行中。
NORMAL_PROGRESS = {
    ('项目立项', '开发中'): ProjectStage.INIT,
    ('研发阶段', '一次送样'): ProjectStage.RND,
    ('客户小试', '合格待产'): ProjectStage.PILOT,
    ('客户中试', '已中试'): ProjectStage.MID_TEST,
    ('客户中试', '小批量试产'): ProjectStage.MID_TEST,
    ('客户量产下单', '已量产'): ProjectStage.MASS_PROD,
    ('客户量产下单', '量产'): ProjectStage.MASS_PROD,
}
STAGE_SEQUENCE = (
    ProjectStage.INIT,
    ProjectStage.RND,
    ProjectStage.PILOT,
    ProjectStage.MID_TEST,
    ProjectStage.MASS_PROD,
    ProjectStage.ORDER,
)
DESCRIPTION_PARTS = (
    ('项目描述：异常原因', '异常原因'),
    ('项目描述：备注', '备注'),
    ('项目描述：不合格原因', '不合格原因'),
    ('项目描述：总出货数量（吨）', '总出货数量（吨）'),
)
EXPORT_HEADERS = (
    '项目编号', '项目名称', '项目等级', '业务板块', '研发负责人', '项目描述',
    '立项时间', '当前阶段', '当前节点备注', '进度百分比', '是否终止',
    '客户编码', '客户公司全称', '终端主机厂', '业务负责人',
    '制件名称', '产品代码/零件号', '制件用途', '制件颜色',
    '研发高可落地性项目', '业务高可落地性项目',
    '目标成本 (元/kg)', '竞品售价 (元/kg)', '预估市场订单用量 (吨/年)',
    '第一次客户送样时间', '第一次客户小试时间', '第一次小试完成周期 (天)',
    '中试进行时间', '量产进行时间', '档案最后更新',
    '项目协同成员', '销售协同成员',
    '成品SAP码', '材料牌号', '材料类型', '进度节点',
)
SYNC_CELL_LIMIT = 500


class RowError(Exception):
    """写入「同步情况」的行级错误。数据库改动由外层事务回滚。"""


def _text(value):
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'TRUE' if value else ''
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(value, 'g')
    if isinstance(value, int):
        return str(value)
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d %H:%M:%S')
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _as_bool(value):
    if isinstance(value, bool):
        return value
    if value is None or value == '':
        return False
    if isinstance(value, (int, float)):
        return bool(value)
    return _text(value).lower() in {'1', 'true', 'yes', 'y', '是', 'x'}


def _as_datetime(value):
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        dt = value
    elif isinstance(value, date):
        dt = datetime(value.year, value.month, value.day)
    else:
        raw = _text(value)
        if not raw:
            return None
        parsed = None
        for fmt, size in (('%Y-%m-%d %H:%M:%S', 19), ('%Y-%m-%d', 10), ('%Y/%m/%d', 10)):
            try:
                parsed = datetime.strptime(raw[:size], fmt)
                break
            except ValueError:
                continue
        if parsed is None:
            raise RowError(f'无法解析立项时间: {value!r}')
        dt = parsed
    if timezone.is_naive(dt):
        dt = timezone.make_aware(dt, timezone.get_current_timezone())
    return dt


def _fit(text, limit):
    if len(text) <= limit:
        return text, False
    return text[:limit], True


def _person_label(user):
    if user is None:
        return ''
    full = (user.get_full_name() or '').strip()
    return full or user.username


def _member_names(members):
    ordered = sorted(members, key=lambda member: member.pk)
    return '、'.join(label for label in (_person_label(member.user) for member in ordered) if label)


def _decimal_text(value):
    if value is None:
        return ''
    return format(value, 'f')


def _date_text(value):
    if value is None:
        return ''
    return value.isoformat()


def _current_node(project):
    """与当前阶段同一口径：终止节点优先，否则取第一个未结束节点。"""
    nodes = sorted(project.nodes.all(), key=lambda node: (node.order, node.pk))
    if not nodes:
        return None
    terminated = [node for node in nodes if node.status == 'TERMINATED']
    if terminated:
        return terminated[-1]
    for node in nodes:
        if node.status not in {'DONE', 'FAILED', 'FEEDBACK'}:
            return node
    return nodes[-1]


def _user_keys(user):
    keys = set()
    for raw in (user.first_name, user.last_name, user.username, user.get_full_name()):
        key = (raw or '').replace(' ', '').strip()
        if key:
            keys.add(key)
    return keys


class UserIndex:
    def __init__(self):
        self._by_key = {}
        for user in User.objects.all().only('id', 'first_name', 'last_name', 'username'):
            for key in _user_keys(user):
                self._by_key.setdefault(key, []).append(user)

    def lookup(self, name):
        """返回 (user 或 None, 'empty'|'ok'|'missing'|'ambiguous')。"""
        key = (name or '').replace(' ', '').strip()
        if not key:
            return None, 'empty'
        matches = {user.pk: user for user in self._by_key.get(key, [])}
        if not matches:
            return None, 'missing'
        if len(matches) > 1:
            return None, 'ambiguous'
        return next(iter(matches.values())), 'ok'


def _description(row):
    parts = []
    for column, label in DESCRIPTION_PARTS:
        text = row.get(column, '')
        if text:
            parts.append(f'{label}：{text}')
    return '\n\n'.join(parts)


def _round1(project, stage):
    try:
        return project.nodes.get(stage=stage, round=1)
    except ProjectNode.DoesNotExist as exc:
        raise RowError(f'缺少第 1 轮节点 {stage}') from exc


def _mark(node, status, remark=None):
    node.status = status
    if remark is not None:
        node.remark = remark
    node.save()


def _complete_through(project, stage):
    target = STAGE_SEQUENCE.index(stage)
    for code in STAGE_SEQUENCE[:target + 1]:
        _mark(_round1(project, code), 'DONE')
    if target + 1 < len(STAGE_SEQUENCE):
        _mark(_round1(project, STAGE_SEQUENCE[target + 1]), 'DOING')


def _iterate_failed_pilot(project, reason):
    """第 1 轮立项、研发完成后，对客户小试走端到端异常迭代。"""
    _mark(_round1(project, ProjectStage.INIT), 'DONE')
    _mark(_round1(project, ProjectStage.RND), 'DONE')
    _round1(project, ProjectStage.PILOT).perform_failure_logic(reason)
    rnd2 = project.nodes.filter(stage=ProjectStage.RND).order_by('-order').first()
    pilot2 = project.nodes.filter(stage=ProjectStage.PILOT).order_by('-order').first()
    if rnd2 is None or pilot2 is None or rnd2.round != 2 or pilot2.round != 2:
        raise RowError('客户小试异常迭代未生成第 2 轮节点')
    return rnd2, pilot2


def apply_progress(project, completed, project_status, fail_reason, exception_reason):
    pair = (completed, project_status)
    if pair == ('暂停', '暂停'):
        _mark(_round1(project, ProjectStage.INIT), 'PAUSED', remark=exception_reason or '')
        return
    if pair == ('终止', '终止'):
        _mark(
            _round1(project, ProjectStage.INIT),
            'TERMINATED',
            remark=exception_reason or '导入：项目终止',
        )
        return
    if pair == ('不合格再开发', '不合格再开发'):
        rnd2, _pilot2 = _iterate_failed_pilot(project, fail_reason or '导入：不合格再开发')
        _mark(rnd2, 'DOING')
        return
    if pair == ('研发阶段', '多次送样'):
        rnd2, pilot2 = _iterate_failed_pilot(project, fail_reason or '导入：多次送样')
        _mark(rnd2, 'DONE')
        _mark(pilot2, 'DOING')
        return
    stage = NORMAL_PROGRESS.get(pair)
    if stage is None:
        if not completed and not project_status:
            raise RowError('项目已完成状态、项目状态都为空')
        raise RowError(f'无法识别的状态组合：已完成={completed or "空"}，项目状态={project_status or "空"}')
    _complete_through(project, stage)


def _progress_done(project):
    init = project.nodes.filter(stage=ProjectStage.INIT, round=1).only('status').first()
    return bool(init and init.status != 'PENDING')


def _cell_rank(cell):
    """「已同步」以及「已同步-异常1；异常2」都算该步已经写过。"""
    for rank in (4, 3, 2, 1):
        mark = RANK_MARK[rank]
        if cell == mark or cell.startswith(mark + '-'):
            return rank
    return 0


def _rank(cell, project):
    rank = _cell_rank(cell)
    if project is not None:
        rank = max(rank, 1)
        if _progress_done(project):
            rank = 4
    return rank


def format_sync_status(mark, issues):
    """无异常只写阶段标记；有异常写成「已同步-异常1；异常2」。"""
    unique = list(dict.fromkeys(issue for issue in issues if issue))
    if not unique:
        return mark
    text = mark + '-' + '；'.join(unique)
    if len(text) <= SYNC_CELL_LIMIT:
        return text
    return text[:SYNC_CELL_LIMIT - 1] + '…'


def _created_at_issue(value):
    if value is None or value == '':
        return '立项时间为空，未写入立项时间'
    if not isinstance(value, (datetime, date)) and not _text(value):
        return '立项时间为空，未写入立项时间'
    try:
        parsed = _as_datetime(value)
    except RowError:
        return '立项时间无法解析'
    if parsed is None:
        return '立项时间为空，未写入立项时间'
    return ''


def _project_issues(project, row):
    issues = []
    name_limit = Project._meta.get_field('name').max_length
    if len(row['项目名称']) > name_limit:
        issues.append('项目名称超长已截断')
    grade_name = row['项目等级']
    if grade_name and not project.grade_id:
        issues.append(f'项目等级「{grade_name}」无法关联，已留空')
    segment_name = row['*业务板块']
    if not segment_name:
        issues.append('业务板块为空，未关联')
    elif not project.business_segment_id:
        issues.append(f'业务板块「{segment_name}」无法关联，已留空')
    created_issue = _created_at_issue(row['立项时间'])
    if created_issue:
        issues.append(created_issue)
    return issues


def _repository_issues(project, row, context):
    issues = []
    repo = ProjectRepository.objects.filter(project_id=project.pk).only('customer_id', 'salesperson_id').first()
    if repo is None:
        return ['项目档案不存在']
    code = row['*客户编码']
    if not code:
        issues.append('客户编码为空，未关联客户')
    elif code.upper() in {'#N/A', 'N/A'}:
        issues.append(f'客户编码「{code}」无效，未关联客户')
    elif context['customers'].get(code) is None or repo.customer_id is None:
        issues.append(f'客户编码「{code}」无法关联，未关联客户')

    salesperson_name = row['*业务负责人']
    _matched, state = context['users'].lookup(salesperson_name)
    if state == 'empty':
        issues.append('业务负责人为空，未关联业务员')
    elif state == 'ambiguous':
        issues.append(f'业务负责人「{salesperson_name}」匹配到多名用户，未关联')
    elif state == 'missing' or repo.salesperson_id is None:
        issues.append(f'业务负责人「{salesperson_name}」无法关联')

    too_long = False
    for column, field_name in (
        ('制件名称(零件名或总成)', 'product_name'),
        ('制件用途（用在哪个产品）', 'part_usage'),
        ('制件颜色（项目档案）', 'part_color'),
    ):
        limit = ProjectRepository._meta.get_field(field_name).max_length
        if len(row[column]) > limit:
            too_long = True
    if too_long:
        issues.append('制件名称、用途或颜色超长已截断')
    return issues


def _material_issues(project, row, context):
    issues = []
    sap = row['*成品SAP码']
    type_name = row['*材料类型']
    if not sap:
        if type_name:
            issues.append(f'成品SAP码为空，材料类型「{type_name}」未关联成品')
        else:
            issues.append('成品SAP码为空，未关联成品')
            issues.append('材料类型为空，未关联成品')
        return issues

    material = None
    if project.material_id:
        material = MaterialLibrary.objects.select_related('category').filter(pk=project.material_id).first()
    if material is None:
        if not type_name:
            issues.append(f'材料类型为空，无法创建成品 {sap}')
        elif context['types'].get(type_name) is None:
            issues.append(f'材料类型「{type_name}」无法关联，成品 {sap} 未创建')
        elif row['材料牌号'] and MaterialLibrary.objects.filter(grade_name=row['材料牌号']).exists():
            issues.append(f'材料牌号「{row["材料牌号"]}」已被占用，成品 {sap} 未创建')
        else:
            issues.append(f'成品 {sap} 未关联')
        return issues

    if not type_name:
        issues.append(f'材料类型为空，已挂接成品「{material.grade_name}」')
    elif material.category_id and material.category.name != type_name:
        issues.append(
            f'材料类型「{type_name}」无法关联：成品现有类型为「{material.category.name}」，已挂接但不修改'
        )
    return issues


def collect_issues(project, row, rank, context):
    """按已经跑完的步骤，收集仍未关联上的项。"""
    if project is None or rank <= 0:
        return []
    issues = _project_issues(project, row)
    if rank >= 2:
        issues.extend(_repository_issues(project, row, context))
    if rank >= 3:
        issues.extend(_material_issues(project, row, context))
    return issues


def _new_grade_name():
    for _ in range(20):
        suffix = ''.join(secrets.choice(GRADE_ALPHABET) for _ in range(6))
        name = f'{GRADE_PREFIX}{suffix}'
        if not MaterialLibrary.objects.filter(grade_name=name).exists():
            return name
    raise RowError('生成材料牌号多次冲突，请重试')


def _read_row(worksheet, excel_row, columns):
    row = {}
    for name, col in columns.items():
        row[name] = worksheet.cell(excel_row, col).value
    return row


def _blank_row(row):
    return not any(_text(row.get(name)) for name in HEADERS if name != '同步情况')


class Command(BaseCommand):
    help = '从 init/项目导入.xlsx 按项目编号分步同步项目，或导出全部项目'

    def add_arguments(self, parser):
        parser.add_argument('--file', dest='file_path', default='', help='导入 Excel，默认 init/项目导入.xlsx')
        parser.add_argument(
            '--step', action='append', choices=STEP_ORDER, dest='steps',
            help='只执行该步骤，可重复。缺省为 project、repository、material、progress',
        )
        parser.add_argument('--limit', type=int, default=0, help='本轮最多处理 N 行，0 表示不限制')
        parser.add_argument('--dry-run', action='store_true', help='只演练，不写数据库，也不改 Excel')
        parser.add_argument(
            '--export', nargs='?', const='', default=None,
            help='导出全部项目到 xlsx。不带路径时写入 init/项目导出.xlsx',
        )

    def handle(self, *args, **options):
        self._mute_notifications()
        try:
            if options['export'] is not None:
                self._export(options['export'])
                return
            self._import(options)
        finally:
            self._unmute_notifications()

    def _mute_notifications(self):
        # 注册时 sender 是 ProjectNode。只按 dispatch_uid、不带 sender 拆不掉。
        post_save.disconnect(sender=ProjectNode, dispatch_uid='notif_project.node_updated')

    def _unmute_notifications(self):
        from app_project.notifications import _register_project_bindings
        _register_project_bindings()

    def _import(self, options):
        if options['limit'] < 0:
            raise CommandError('--limit 不能为负数')
        file_path = options['file_path'] or os.path.join(settings.BASE_DIR, 'init', '项目导入.xlsx')
        if not os.path.exists(file_path):
            raise CommandError(f'找不到数据文件: {file_path}')
        try:
            import openpyxl
        except ImportError as exc:
            raise CommandError('缺少 openpyxl，无法读取 xlsx') from exc

        steps = [step for step in STEP_ORDER if step in (options['steps'] or STEP_ORDER)]
        workbook = openpyxl.load_workbook(file_path)
        try:
            worksheet = workbook.active
            columns = self._header_map(worksheet)
            self.stdout.write(self.style.MIGRATE_HEADING(
                f"开始同步项目: {file_path}  步骤: {', '.join(steps)}"
            ))
            if options['dry_run']:
                self.stdout.write(self.style.WARNING('dry-run：不会写数据库，也不会改 Excel'))
            context = self._load_context()
            if options['dry_run']:
                with transaction.atomic():
                    counts = self._walk(worksheet, columns, steps, options, context, workbook, file_path, commit=False)
                    transaction.set_rollback(True)
            else:
                counts = self._walk(worksheet, columns, steps, options, context, workbook, file_path, commit=True)
        finally:
            workbook.close()
        self._report(counts)

    def _header_map(self, worksheet):
        columns = {}
        for col in range(1, worksheet.max_column + 1):
            name = _text(worksheet.cell(1, col).value)
            if name and name not in columns:
                columns[name] = col
        missing = [name for name in HEADERS if name not in columns]
        if missing:
            raise CommandError('表头缺少列: ' + '、'.join(missing))
        return columns

    def _load_context(self):
        return {
            'users': UserIndex(),
            'grades': {item.name: item for item in GradeFactor.objects.all()},
            'segments': {item.name: item for item in BusinessSegment.objects.all()},
            'types': {item.name: item for item in MaterialType.objects.all()},
            'customers': {
                item.customer_code: item
                for item in Customer.objects.exclude(customer_code__isnull=True).exclude(customer_code='')
            },
            'workflow': ProjectConfig.get().default_approval_workflow,
        }

    def _walk(self, worksheet, columns, steps, options, context, workbook, file_path, commit):
        counts = {'done': 0, 'skipped': 0, 'failed': 0}
        limit = options['limit']
        processed = 0
        for excel_row in range(2, worksheet.max_row + 1):
            if limit and processed >= limit:
                break
            raw = _read_row(worksheet, excel_row, columns)
            if _blank_row(raw):
                continue
            row = {name: _text(raw.get(name)) for name in HEADERS}
            row['立项时间'] = raw.get('立项时间')
            row['研发高可落地性项目'] = _as_bool(raw.get('研发高可落地性项目'))
            row['业务高可落地性项目'] = _as_bool(raw.get('业务高可落地性项目'))
            code = row['项目编号']
            cell = row['同步情况']
            if not code:
                self._fail(worksheet, columns, excel_row, '', RowError('项目编号为空'), counts, workbook, file_path, commit)
                processed += 1
                continue
            project = Project.objects.filter(code=code).first()
            rank = _rank(cell, project)
            pending = [step for step in steps if STEP_RANK[step] > rank]
            if not pending:
                grade_name, failed = self._retry_unlinked_material(
                    project, row, steps, context, counts, worksheet, columns, excel_row, code,
                    workbook, file_path, commit,
                )
                if failed:
                    continue
                self._refresh_status(
                    worksheet, columns, excel_row, cell, rank, project, row, context,
                    commit, workbook, file_path, code, grade_name,
                )
                counts['skipped'] += 1
                continue
            processed += 1
            try:
                with transaction.atomic():
                    mark, grade_name, issues = self._apply_row(project, row, pending, rank, context)
            except RowError as exc:
                self._fail(worksheet, columns, excel_row, code, exc, counts, workbook, file_path, commit)
                continue
            except Exception as exc:
                self._fail(worksheet, columns, excel_row, code, exc, counts, workbook, file_path, commit)
                continue
            for issue in issues:
                self.stdout.write(self.style.WARNING(f'  [!] 第 {excel_row} 行 {code}: {issue}'))
            if commit:
                self._write_success(worksheet, columns, excel_row, mark, grade_name)
                self._save_workbook(workbook, file_path)
            self.stdout.write(f'  [+] 第 {excel_row} 行 {code}: {mark}' + ('（dry-run）' if not commit else ''))
            counts['done'] += 1
        return counts

    def _retry_unlinked_material(
        self, project, row, steps, context, counts, worksheet, columns, excel_row, code,
        workbook, file_path, commit,
    ):
        """进度已写入后材料类型才补上时，再挂一次成品，不再推进节点。"""
        if not commit or project is None or 'material' not in steps:
            return None, False
        if project.material_id or not row['*成品SAP码']:
            return None, False
        try:
            with transaction.atomic():
                grade_name = self._link_material(project, row, context)
        except RowError as exc:
            self._fail(worksheet, columns, excel_row, code, exc, counts, workbook, file_path, commit)
            return None, True
        except Exception as exc:
            self._fail(worksheet, columns, excel_row, code, exc, counts, workbook, file_path, commit)
            return None, True
        project.refresh_from_db()
        return grade_name, False

    def _refresh_status(
        self, worksheet, columns, excel_row, cell, rank, project, row, context,
        commit, workbook, file_path, code, grade_name,
    ):
        if not commit or rank <= 0 or project is None:
            return
        mark = format_sync_status(RANK_MARK[rank], collect_issues(project, row, rank, context))
        changed = False
        if cell != mark:
            worksheet.cell(excel_row, columns['同步情况'], mark)
            changed = True
        if grade_name:
            worksheet.cell(excel_row, columns['材料牌号'], grade_name)
            changed = True
        if not changed:
            return
        self._save_workbook(workbook, file_path)
        self.stdout.write(f'  [=] 第 {excel_row} 行 {code}: {mark}')

    def _fail(self, worksheet, columns, excel_row, code, exc, counts, workbook, file_path, commit):
        counts['failed'] += 1
        message = f'{exc.__class__.__name__}: {exc}'
        self.stdout.write(self.style.ERROR(f'  [x] 第 {excel_row} 行 {code}: {message}'))
        if not commit:
            return
        worksheet.cell(excel_row, columns['同步情况'], message[:SYNC_CELL_LIMIT])
        self._save_workbook(workbook, file_path)

    def _write_success(self, worksheet, columns, excel_row, mark, grade_name):
        worksheet.cell(excel_row, columns['同步情况'], mark)
        if grade_name:
            worksheet.cell(excel_row, columns['材料牌号'], grade_name)

    def _save_workbook(self, workbook, file_path):
        try:
            workbook.save(file_path)
        except OSError as exc:
            raise CommandError(f'数据库已写入，但保存 Excel 失败（请先关闭文件）: {exc}') from exc

    def _apply_row(self, project, row, pending, rank, context):
        grade_name = None
        if 'project' in pending:
            project = self._create_project(row, context)
        if project is None:
            raise RowError('项目尚未创建，请先执行 --step project')
        if 'repository' in pending:
            self._fill_repository(project, row, context)
        if 'material' in pending:
            grade_name = self._link_material(project, row, context)
        if 'progress' in pending and not _progress_done(project):
            apply_progress(
                project,
                row['*项目已完成状态'],
                row['项目状态'],
                row['项目描述：不合格原因'],
                row['项目描述：异常原因'],
            )
        new_rank = max([rank, *(STEP_RANK[step] for step in pending)])
        issues = collect_issues(project, row, new_rank, context)
        return format_sync_status(RANK_MARK[new_rank], issues), grade_name, issues

    def _create_project(self, row, context):
        name = row['项目名称']
        if not name:
            raise RowError('项目名称为空')
        name, _clipped = _fit(name, Project._meta.get_field('name').max_length)
        manager, state = context['users'].lookup(row['*研发负责人'])
        if state == 'empty':
            raise RowError('研发负责人为空')
        if state == 'missing':
            raise RowError(f"找不到研发负责人「{row['*研发负责人']}」")
        if state == 'ambiguous':
            raise RowError(f"研发负责人「{row['*研发负责人']}」匹配到多名用户")

        grade_name = row['项目等级']
        grade = context['grades'].get(grade_name) if grade_name else None
        segment_name = row['*业务板块']
        segment = context['segments'].get(segment_name) if segment_name else None

        project = Project.objects.create(
            code=row['项目编号'],
            name=name,
            grade=grade,
            business_segment=segment,
            description=_description(row),
            manager=manager,
            approval_workflow=context['workflow'],
        )
        created_at = _as_datetime(row['立项时间'])
        if created_at is not None:
            Project.objects.filter(pk=project.pk).update(created_at=created_at)
        return project

    def _fill_repository(self, project, row, context):
        repo, _created = ProjectRepository.objects.get_or_create(project=project)
        code = row['*客户编码']
        if not code or code.upper() in {'#N/A', 'N/A'}:
            repo.customer = None
        else:
            repo.customer = context['customers'].get(code)

        salesperson, state = context['users'].lookup(row['*业务负责人'])
        repo.salesperson = salesperson if state == 'ok' else None

        product, _clipped = _fit(row['制件名称(零件名或总成)'], 100)
        usage, _usage_clipped = _fit(row['制件用途（用在哪个产品）'], 100)
        color, _color_clipped = _fit(row['制件颜色（项目档案）'], 50)
        repo.product_name = product
        repo.part_usage = usage
        repo.part_color = color
        repo.rd_high_feasibility = row['研发高可落地性项目']
        repo.biz_high_feasibility = row['业务高可落地性项目']
        repo.save()

    def _link_material(self, project, row, context):
        sap = row['*成品SAP码']
        if not sap:
            return None
        type_name = row['*材料类型']
        material = (
            MaterialLibrary.objects.filter(sap_material_code=sap)
            .select_related('category')
            .first()
        )
        if material is not None:
            if project.material_id != material.pk:
                project.material = material
                project.save(update_fields=['material'])
            return None

        material_type = context['types'].get(type_name) if type_name else None
        if material_type is None:
            return None

        supplied = row['材料牌号']
        generated = False
        if supplied:
            if MaterialLibrary.objects.filter(grade_name=supplied).exists():
                return None
            grade_name = supplied
        else:
            grade_name = _new_grade_name()
            generated = True
        material = MaterialLibrary.objects.create(
            grade_name=grade_name,
            sap_material_code=sap,
            category=material_type,
            creator=project.manager,
            is_published=False,
        )
        project.material = material
        project.save(update_fields=['material'])
        return grade_name if generated else None

    def _report(self, counts):
        self.stdout.write(self.style.MIGRATE_HEADING('\n--- 同步完成 ---'))
        self.stdout.write(f"处理: {counts['done']}")
        self.stdout.write(f"跳过: {counts['skipped']}")
        if counts['failed']:
            self.stdout.write(self.style.ERROR(f"失败: {counts['failed']}"))
        else:
            self.stdout.write(f"失败: {counts['failed']}")

    def _export(self, raw_path):
        try:
            import openpyxl
        except ImportError as exc:
            raise CommandError('缺少 openpyxl，无法写入 xlsx') from exc
        file_path = raw_path or os.path.join(settings.BASE_DIR, 'init', '项目导出.xlsx')
        parent = os.path.dirname(file_path)
        if parent:
            os.makedirs(parent, exist_ok=True)
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = '项目'
        worksheet.append(EXPORT_HEADERS)
        projects = Project.objects.select_related(
            'manager', 'grade', 'business_segment', 'material', 'material__category',
            'repository', 'repository__customer', 'repository__salesperson', 'repository__oem',
        ).prefetch_related('nodes', 'members__user', 'sales_members__user').order_by('code')
        count = 0
        for project in projects:
            worksheet.append(self._export_row(project))
            count += 1
        workbook.save(file_path)
        workbook.close()
        self.stdout.write(self.style.SUCCESS(f'已导出 {count} 个项目: {file_path}'))

    def _export_row(self, project):
        repo = getattr(project, 'repository', None)
        material = project.material
        customer = repo.customer if repo and repo.customer_id else None
        oem = repo.oem if repo and repo.oem_id else None
        created = timezone.localtime(project.created_at).strftime('%Y-%m-%d %H:%M:%S')
        updated = timezone.localtime(repo.updated_at).strftime('%Y-%m-%d %H:%M:%S') if repo else ''
        nodes = project.nodes.all()
        current = _current_node(project)
        node_text = '\n'.join(
            f'{node.get_stage_display()} 第{node.round}轮 {node.get_status_display()}'
            for node in nodes
        )
        return [
            project.code,
            project.name,
            project.grade.name if project.grade_id else '',
            project.business_segment.name if project.business_segment_id else '',
            _person_label(project.manager),
            project.description,
            created,
            project.get_current_stage_display(),
            (current.remark or '') if current else '',
            project.progress_percent,
            '是' if project.is_terminated else '否',
            customer.customer_code if customer else '',
            customer.company_name if customer else '',
            oem.name if oem else '',
            _person_label(repo.salesperson) if repo else '',
            repo.product_name if repo else '',
            repo.product_code if repo else '',
            repo.part_usage if repo else '',
            repo.part_color if repo else '',
            '是' if repo and repo.rd_high_feasibility else '否',
            '是' if repo and repo.biz_high_feasibility else '否',
            _decimal_text(repo.target_cost) if repo else '',
            _decimal_text(repo.competitor_price) if repo else '',
            _decimal_text(repo.estimated_order_volume) if repo else '',
            _date_text(repo.first_sample_date) if repo else '',
            _date_text(repo.first_trial_date) if repo else '',
            repo.first_trial_cycle_days if repo and repo.first_trial_cycle_days is not None else '',
            _date_text(repo.pilot_date) if repo else '',
            _date_text(repo.mass_production_date) if repo else '',
            updated,
            _member_names(project.members.all()),
            _member_names(project.sales_members.all()),
            material.sap_material_code if material else '',
            material.grade_name if material else '',
            material.category.name if material else '',
            node_text,
        ]
