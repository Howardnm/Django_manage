import os
import re
import shutil
import tempfile
from datetime import date, datetime
from decimal import Decimal
from io import StringIO

import openpyxl
from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.test import TestCase
from django.utils import timezone

from app_material.models import MaterialLibrary, MaterialType
from app_notification.models import Notification
from app_project.management.commands.import_projects import GRADE_PREFIX, HEADERS
from app_project.models import (
    BusinessSegment, Project, ProjectMember, ProjectNode, ProjectSalesMember, ProjectStage,
)
from app_repository.models import Customer, GradeFactor, OEM

User = get_user_model()


def _user(username, first_name):
    user = User.objects.create_user(username=username, email=f'{username}@test.local', password='x')
    user.first_name = first_name
    user.save(update_fields=['first_name'])
    return user


def _base_row(**overrides):
    data = {name: None for name in HEADERS}
    data.update({
        '项目编号': '26B02001001',
        '项目名称': '导入测试项目',
        '项目等级': 'A',
        '*材料类型': 'PP',
        '*业务板块': '汽配',
        '*研发负责人': '彭万',
        '立项时间': datetime(2026, 1, 5, 0, 0, 0),
        '*项目已完成状态': '研发阶段',
        '项目状态': '一次送样',
        '项目描述：异常原因': '外观偏差',
        '项目描述：备注': '首次送样',
        '项目描述：不合格原因': '强度不够',
        '项目描述：总出货数量（吨）': 3.5,
        '*客户编码': '10200209',
        '*业务负责人': '冯丹龙',
        '制件名称(零件名或总成)': '风叶',
        '制件用途（用在哪个产品）': '台扇',
        '制件颜色（项目档案）': '本色',
        '研发高可落地性项目': True,
        '业务高可落地性项目': False,
    })
    data.update(overrides)
    return [data[name] for name in HEADERS]


def _write(path, rows):
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.append(list(HEADERS))
    for row in rows:
        worksheet.append(row)
    workbook.save(path)
    workbook.close()


def _cell(path, row, column):
    workbook = openpyxl.load_workbook(path, data_only=True)
    value = workbook.active.cell(row, column).value
    workbook.close()
    return value


def _nodes(project):
    return list(project.nodes.order_by('order').values_list('stage', 'round', 'status'))


class ImportProjectsTests(TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)
        self.path = os.path.join(self.tmp, 'import.xlsx')
        self.rd = _user('rd_peng', '彭万')
        self.sales = _user('sales_feng', '冯丹龙')
        Customer.objects.create(customer_code='10200209', company_name='测试客户')
        MaterialType.objects.create(name='PP')
        GradeFactor.objects.create(name='A', factor='1.20')
        BusinessSegment.objects.create(name='汽配')

    def _run(self, **options):
        out = StringIO()
        call_command('import_projects', file_path=self.path, stdout=out, **options)
        return out.getvalue()

    def test_first_sample_links_archive_and_writes_synced(self):
        _write(self.path, [_base_row()])
        self._run()

        project = Project.objects.get(code='26B02001001')
        self.assertEqual(project.manager_id, self.rd.pk)
        self.assertEqual(project.grade.name, 'A')
        self.assertEqual(project.business_segment.name, '汽配')
        self.assertIn('异常原因：外观偏差', project.description)
        self.assertIn('总出货数量（吨）：3.5', project.description)
        self.assertEqual(timezone.localtime(project.created_at).date().isoformat(), '2026-01-05')
        self.assertEqual(Notification.objects.filter(type='project.node_updated').count(), 0)
        repo = project.repository
        self.assertEqual(repo.customer.customer_code, '10200209')
        self.assertEqual(repo.salesperson_id, self.sales.pk)
        self.assertEqual(repo.product_name, '风叶')
        self.assertEqual(repo.part_usage, '台扇')
        self.assertEqual(repo.part_color, '本色')
        self.assertTrue(repo.rd_high_feasibility)
        self.assertFalse(repo.biz_high_feasibility)
        self.assertEqual(_nodes(project), [
            (ProjectStage.INIT, 1, 'DONE'),
            (ProjectStage.RND, 1, 'DONE'),
            (ProjectStage.PILOT, 1, 'DOING'),
            (ProjectStage.MID_TEST, 1, 'PENDING'),
            (ProjectStage.MASS_PROD, 1, 'PENDING'),
            (ProjectStage.ORDER, 1, 'PENDING'),
            (ProjectStage.MASS_TRACK, 1, 'PENDING'),
        ])
        self.assertEqual(_cell(self.path, 2, 1), '已同步-成品SAP码为空，材料类型「PP」未关联成品')
        self.assertIsNone(project.material_id)

        self._run()
        self.assertEqual(Project.objects.filter(code='26B02001001').count(), 1)
        self.assertEqual(project.nodes.count(), 7)

    def test_multiple_samples_and_redevelopment(self):
        _write(self.path, [
            _base_row(项目编号='MULTI0001', **{'项目状态': '多次送样'}),
            _base_row(
                项目编号='REDEV0001',
                **{'*项目已完成状态': '不合格再开发', '项目状态': '不合格再开发'},
            ),
        ])
        self._run()

        multi = Project.objects.get(code='MULTI0001')
        self.assertEqual(_nodes(multi), [
            (ProjectStage.INIT, 1, 'DONE'),
            (ProjectStage.RND, 1, 'DONE'),
            (ProjectStage.PILOT, 1, 'FAILED'),
            (ProjectStage.RND, 2, 'DONE'),
            (ProjectStage.PILOT, 2, 'DOING'),
            (ProjectStage.MID_TEST, 1, 'PENDING'),
            (ProjectStage.MASS_PROD, 1, 'PENDING'),
            (ProjectStage.ORDER, 1, 'PENDING'),
            (ProjectStage.MASS_TRACK, 1, 'PENDING'),
        ])
        failed = multi.nodes.get(stage=ProjectStage.PILOT, round=1)
        self.assertEqual(failed.remark, '强度不够')

        redev = Project.objects.get(code='REDEV0001')
        self.assertEqual(
            list(redev.nodes.filter(stage=ProjectStage.RND).order_by('round').values_list('round', 'status')),
            [(1, 'DONE'), (2, 'DOING')],
        )
        self.assertEqual(
            list(redev.nodes.filter(stage=ProjectStage.PILOT).order_by('round').values_list('round', 'status')),
            [(1, 'FAILED'), (2, 'PENDING')],
        )

        before = ProjectNode.objects.filter(project=multi).count()
        self._run()
        self.assertEqual(ProjectNode.objects.filter(project=multi).count(), before)

    def test_pause_and_terminate_keep_pipeline(self):
        _write(self.path, [
            _base_row(项目编号='PAUSE0001', **{'*项目已完成状态': '暂停', '项目状态': '暂停'}),
            _base_row(项目编号='STOP00001', **{'*项目已完成状态': '终止', '项目状态': '终止'}),
        ])
        self._run()

        paused = Project.objects.get(code='PAUSE0001')
        paused.refresh_from_db()
        self.assertEqual(paused.nodes.get(stage=ProjectStage.INIT).status, 'PAUSED')
        self.assertEqual(paused.nodes.get(stage=ProjectStage.INIT).remark, '外观偏差')
        self.assertFalse(paused.is_terminated)
        self.assertEqual(paused.nodes.count(), 7)

        stopped = Project.objects.get(code='STOP00001')
        stopped.refresh_from_db()
        self.assertTrue(stopped.is_terminated)
        self.assertEqual(stopped.nodes.get(stage=ProjectStage.INIT).status, 'TERMINATED')
        self.assertEqual(stopped.nodes.count(), 7)

    def test_mass_production_sets_order_doing(self):
        _write(self.path, [_base_row(
            项目编号='MASS000001',
            **{'*项目已完成状态': '客户量产下单', '项目状态': '已量产'},
        )])
        self._run()
        project = Project.objects.get(code='MASS000001')
        self.assertEqual(project.nodes.get(stage=ProjectStage.MASS_PROD).status, 'DONE')
        self.assertEqual(project.nodes.get(stage=ProjectStage.ORDER).status, 'DOING')
        self.assertEqual(project.nodes.get(stage=ProjectStage.MASS_TRACK).status, 'PENDING')

    def test_empty_grade_name_is_generated_and_written_back(self):
        _write(self.path, [
            _base_row(项目编号='SAP0000001', **{'*成品SAP码': 'A02009000090'}),
            _base_row(项目编号='SAP0000002', **{'*成品SAP码': 'A02009000091'}),
        ])
        self._run()

        names = []
        expected_sap = {'SAP0000001': 'A02009000090', 'SAP0000002': 'A02009000091'}
        for excel_row, code in ((2, 'SAP0000001'), (3, 'SAP0000002')):
            grade = _cell(self.path, excel_row, HEADERS.index('材料牌号') + 1)
            self.assertRegex(grade, rf'^{re.escape(GRADE_PREFIX)}[A-Za-z0-9]{{6}}$')
            material = Project.objects.get(code=code).material
            self.assertEqual(material.grade_name, grade)
            self.assertEqual(material.sap_material_code, expected_sap[code])
            self.assertEqual(material.category.name, 'PP')
            self.assertEqual(material.creator_id, self.rd.pk)
            self.assertFalse(material.is_published)
            names.append(grade)
        self.assertEqual(len(set(names)), 2)
        self.assertEqual(MaterialLibrary.objects.count(), 2)

    def test_existing_sap_is_linked_without_new_grade(self):
        material_type = MaterialType.objects.get(name='PP')
        existing = MaterialLibrary.objects.create(
            grade_name='已有牌号', sap_material_code='A02009000090', category=material_type,
        )
        _write(self.path, [_base_row(项目编号='LINK000001', **{'*成品SAP码': 'A02009000090'})])
        self._run()

        project = Project.objects.get(code='LINK000001')
        self.assertEqual(project.material_id, existing.pk)
        self.assertEqual(MaterialLibrary.objects.count(), 1)
        self.assertIsNone(_cell(self.path, 2, HEADERS.index('材料牌号') + 1))

    def test_missing_material_type_is_recorded_without_rollback(self):
        _write(self.path, [_base_row(**{'*成品SAP码': 'A02009999999', '*材料类型': '不存在的类型'})])
        output = self._run()
        self.assertIn('不存在的类型', output)
        project = Project.objects.get()
        self.assertIsNone(project.material_id)
        self.assertEqual(MaterialLibrary.objects.count(), 0)
        self.assertEqual(
            _cell(self.path, 2, 1),
            '已同步-材料类型「不存在的类型」无法关联，成品 A02009999999 未创建',
        )

        MaterialType.objects.create(name='不存在的类型')
        self._run()
        project.refresh_from_db()
        self.assertEqual(project.material.sap_material_code, 'A02009999999')
        self.assertEqual(project.material.category.name, '不存在的类型')
        self.assertEqual(_cell(self.path, 2, 1), '已同步')
        self.assertRegex(_cell(self.path, 2, HEADERS.index('材料牌号') + 1), rf'^{re.escape(GRADE_PREFIX)}[A-Za-z0-9]{{6}}$')
        self.assertEqual(Project.objects.count(), 1)

    def test_unknown_manager_is_recorded_and_limit_skips_synced_rows(self):
        _write(self.path, [
            _base_row(项目编号='OK00000001'),
            _base_row(项目编号='BADMANAGER', **{'*研发负责人': '查无此人'}),
            _base_row(项目编号='OK00000002'),
        ])
        self._run(limit=1)
        self.assertEqual(list(Project.objects.values_list('code', flat=True)), ['OK00000001'])
        self.assertEqual(_cell(self.path, 2, 1), '已同步-成品SAP码为空，材料类型「PP」未关联成品')
        self.assertIsNone(_cell(self.path, 3, 1))

        self._run(limit=1)
        self.assertIn('查无此人', _cell(self.path, 3, 1))
        self.assertFalse(Project.objects.filter(code='BADMANAGER').exists())
        self.assertIsNone(_cell(self.path, 4, 1))

    def test_dry_run_writes_nothing(self):
        _write(self.path, [_base_row()])
        output = self._run(dry_run=True)
        self.assertIn('dry-run', output)
        self.assertEqual(Project.objects.count(), 0)
        self.assertIsNone(_cell(self.path, 2, 1))

    def test_steps_can_run_separately(self):
        _write(self.path, [_base_row()])
        self._run(steps=['project'])
        project = Project.objects.get(code='26B02001001')
        self.assertEqual(project.nodes.get(stage=ProjectStage.INIT).status, 'PENDING')
        self.assertIsNone(project.repository.customer_id)
        self.assertEqual(_cell(self.path, 2, 1), '项目已同步')

        self._run(steps=['repository'])
        project.repository.refresh_from_db()
        self.assertEqual(project.repository.customer.customer_code, '10200209')
        self.assertEqual(_cell(self.path, 2, 1), '档案已同步')

    def test_export_includes_project_nodes(self):
        _write(self.path, [_base_row()])
        self._run()
        project = Project.objects.get(code='26B02001001')
        current = project.nodes.get(stage=ProjectStage.PILOT, round=1)
        current.remark = '当前送样备注'
        current.save(update_fields=['remark'])
        export_path = os.path.join(self.tmp, 'export.xlsx')
        call_command('import_projects', export=export_path, stdout=StringIO())
        workbook = openpyxl.load_workbook(export_path, data_only=True)
        worksheet = workbook.active
        headers = [worksheet.cell(1, col).value for col in range(1, worksheet.max_column + 1)]
        self.assertIn('进度节点', headers)
        codes = [worksheet.cell(row, 1).value for row in range(2, worksheet.max_row + 1)]
        self.assertIn('26B02001001', codes)
        row = codes.index('26B02001001') + 2
        node_text = worksheet.cell(row, headers.index('进度节点') + 1).value
        remark = worksheet.cell(row, headers.index('当前节点备注') + 1).value
        workbook.close()
        self.assertIn('已完成', node_text)
        self.assertIn('进行中', node_text)
        self.assertEqual(remark, '当前送样备注')

    def test_export_includes_archive_fields_customer_name_and_members(self):
        _write(self.path, [_base_row()])
        self._run()
        project = Project.objects.get(code='26B02001001')
        oem = OEM.objects.create(name='吉利汽车有限公司', short_name='吉利')
        repo = project.repository
        repo.oem = oem
        repo.product_code = 'PN-001'
        repo.target_cost = Decimal('12.50')
        repo.competitor_price = Decimal('11.00')
        repo.estimated_order_volume = Decimal('80.00')
        repo.first_sample_date = date(2026, 2, 1)
        repo.first_trial_date = date(2026, 3, 1)
        repo.first_trial_cycle_days = 28
        repo.pilot_date = date(2026, 4, 1)
        repo.mass_production_date = date(2026, 6, 1)
        repo.save()
        repo.refresh_from_db()
        mate = _user('mate_li', '李协同')
        sales_mate = _user('sales_wang', '王销售')
        ProjectMember.objects.create(project=project, user=mate, role='RND')
        ProjectMember.objects.create(project=project, user=self.rd, role='LEAD')
        ProjectSalesMember.objects.create(project=project, user=sales_mate, workload_share='100')

        export_path = os.path.join(self.tmp, 'export.xlsx')
        call_command('import_projects', export=export_path, stdout=StringIO())
        workbook = openpyxl.load_workbook(export_path, data_only=True)
        worksheet = workbook.active
        headers = [worksheet.cell(1, col).value for col in range(1, worksheet.max_column + 1)]
        values = {
            headers[col - 1]: worksheet.cell(2, col).value
            for col in range(1, worksheet.max_column + 1)
        }
        workbook.close()

        self.assertEqual(values['客户编码'], '10200209')
        self.assertEqual(values['客户公司全称'], '测试客户')
        self.assertEqual(values['终端主机厂'], '吉利汽车有限公司')
        self.assertEqual(values['产品代码/零件号'], 'PN-001')
        self.assertEqual(values['目标成本 (元/kg)'], '12.50')
        self.assertEqual(values['竞品售价 (元/kg)'], '11.00')
        self.assertEqual(values['预估市场订单用量 (吨/年)'], '80.00')
        self.assertEqual(values['第一次客户送样时间'], '2026-02-01')
        self.assertEqual(values['第一次客户小试时间'], '2026-03-01')
        self.assertEqual(values['第一次小试完成周期 (天)'], 28)
        self.assertEqual(values['中试进行时间'], '2026-04-01')
        self.assertEqual(values['量产进行时间'], '2026-06-01')
        self.assertEqual(
            values['档案最后更新'],
            timezone.localtime(repo.updated_at).strftime('%Y-%m-%d %H:%M:%S'),
        )
        self.assertEqual(values['项目协同成员'], '李协同、彭万')
        self.assertEqual(values['销售协同成员'], '王销售')
        self.assertEqual(values['制件名称'], '风叶')
        self.assertEqual(values['研发高可落地性项目'], '是')
        self.assertEqual(values['业务高可落地性项目'], '否')

    def test_import_does_not_leave_node_notifications_disconnected(self):
        _write(self.path, [_base_row()])
        self._run()
        project = Project.objects.get(code='26B02001001')
        node = project.nodes.get(stage=ProjectStage.PILOT)
        before = Notification.objects.filter(type='project.node_updated').count()
        node.remark = '导入后再更新'
        node.save()
        self.assertGreater(
            Notification.objects.filter(type='project.node_updated', recipient=self.rd).count(),
            before,
        )

    def test_sync_cell_lists_association_issues(self):
        long_name = '项' * 101
        _write(self.path, [
            _base_row(
                项目编号='ISSUE00001',
                项目名称=long_name,
                项目等级='Z',
                立项时间=None,
                **{
                    '*业务板块': '未知板块',
                    '*客户编码': 'NO-SUCH',
                    '*业务负责人': '没有这个人',
                },
            ),
            _base_row(
                项目编号='EMPTY00001',
                **{
                    '*业务板块': '',
                    '*客户编码': '',
                    '*业务负责人': '',
                    '*材料类型': '',
                    '制件颜色（项目档案）': '色' * 51,
                },
            ),
            _base_row(项目编号='NACUST0001', **{'*客户编码': '#N/A'}),
        ])
        self._run()

        self.assertTrue(Project.objects.filter(code='ISSUE00001').exists())
        issue = Project.objects.get(code='ISSUE00001')
        self.assertIsNone(issue.grade_id)
        self.assertIsNone(issue.business_segment_id)
        self.assertIsNone(issue.repository.customer_id)
        self.assertIsNone(issue.repository.salesperson_id)
        self.assertEqual(len(issue.name), 100)
        self.assertEqual(_cell(self.path, 2, 1), '已同步-' + '；'.join([
            '项目名称超长已截断',
            '项目等级「Z」无法关联，已留空',
            '业务板块「未知板块」无法关联，已留空',
            '立项时间为空，未写入立项时间',
            '客户编码「NO-SUCH」无法关联，未关联客户',
            '业务负责人「没有这个人」无法关联',
            '成品SAP码为空，材料类型「PP」未关联成品',
        ]))

        empty = Project.objects.get(code='EMPTY00001')
        self.assertEqual(len(empty.repository.part_color), 50)
        self.assertEqual(_cell(self.path, 3, 1), '已同步-' + '；'.join([
            '业务板块为空，未关联',
            '客户编码为空，未关联客户',
            '业务负责人为空，未关联业务员',
            '制件名称、用途或颜色超长已截断',
            '成品SAP码为空，未关联成品',
            '材料类型为空，未关联成品',
        ]))
        self.assertEqual(
            _cell(self.path, 4, 1),
            '已同步-客户编码「#N/A」无效，未关联客户；成品SAP码为空，材料类型「PP」未关联成品',
        )

    def test_ambiguous_salesperson_is_listed(self):
        _user('sales_feng_2', '冯丹龙')
        _write(self.path, [_base_row(项目编号='AMBIG00001')])
        self._run()
        project = Project.objects.get(code='AMBIG00001')
        self.assertIsNone(project.repository.salesperson_id)
        self.assertEqual(
            _cell(self.path, 2, 1),
            '已同步-业务负责人「冯丹龙」匹配到多名用户，未关联；成品SAP码为空，材料类型「PP」未关联成品',
        )

    def test_material_conflicts_are_listed_without_rollback(self):
        other = MaterialType.objects.create(name='ABS')
        existing = MaterialLibrary.objects.create(
            grade_name='已有牌号', sap_material_code='A02009000090', category=other,
        )
        MaterialLibrary.objects.create(
            grade_name='另一条', sap_material_code='A02009000090', category=other,
        )
        MaterialLibrary.objects.create(
            grade_name='占用牌号', sap_material_code='A02009000091', category=other,
        )
        _write(self.path, [
            _base_row(项目编号='MISMATCH01', **{'*成品SAP码': 'A02009000090'}),
            _base_row(项目编号='OCCUPIED01', **{'*成品SAP码': 'A02009999998', '材料牌号': '占用牌号'}),
        ])
        self._run()

        linked = Project.objects.get(code='MISMATCH01')
        self.assertEqual(linked.material_id, existing.pk)
        self.assertEqual(_cell(self.path, 2, 1), '已同步-' + '；'.join([
            'SAP「A02009000090」对应 2 条成品，已选用「已有牌号」',
            '材料类型「PP」无法关联：成品现有类型为「ABS」，已挂接但不修改',
        ]))

        occupied = Project.objects.get(code='OCCUPIED01')
        self.assertIsNone(occupied.material_id)
        self.assertEqual(
            _cell(self.path, 3, 1),
            '已同步-材料牌号「占用牌号」已被占用，成品 A02009999998 未创建',
        )
